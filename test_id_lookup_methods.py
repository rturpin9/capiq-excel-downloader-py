"""
Test three ID lookup methods side by side:
  1. CIQ()          - current builder approach
  2. CIQRANGEA()    - legacy approach (known broken for some things in Pro compat)
  3. SPG()          - Pro native approach

Tests with identifiers that CIQ() can't resolve: company names and CUSIPs.
"""
import os
import time
import openpyxl
import pythoncom

from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

pythoncom.CoInitialize()

# Identifiers that failed with CIQ(), plus some that worked for comparison
TEST_IDS = [
    ("NVDA", "ticker"),
    ("Microsoft", "name"),
    ("Microsoft Corporation", "full name"),
    ("03783310", "CUSIP (Apple)"),
    ("US0378331005", "ISIN (Apple)"),
    ("TSX:CSU", "exchange:ticker"),
]

XLSX_PATH = r"C:\Claude\test_id_methods.xlsx"


def create_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Methods"

    # Headers
    headers = ["Input", "Type",
               "CIQ() ID", "CIQ() Name",
               "CIQRANGEA() ID", "CIQRANGEA() Name",
               "SPG() ID", "SPG() Name"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    for i, (identifier, id_type) in enumerate(TEST_IDS):
        row = i + 2
        ws.cell(row=row, column=1, value=identifier)
        ws.cell(row=row, column=2, value=id_type)

        # Method 1: CIQ()
        ws.cell(row=row, column=3,
                value=f'=CIQ("{identifier}","IQ_COMPANY_ID")')
        ws.cell(row=row, column=4,
                value=f'=CIQ("{identifier}","IQ_COMPANY_NAME")')

        # Method 2: CIQRANGEA() with QUICK_MATCH
        ws.cell(row=row, column=5,
                value=f'=CIQRANGEA("{identifier}","IQ_COMPANY_ID_QUICK_MATCH",1,1)')
        ws.cell(row=row, column=6,
                value=f'=CIQRANGEA("{identifier}","IQ_COMPANY_NAME_QUICK_MATCH",1,1)')

        # Method 3: SPG()
        ws.cell(row=row, column=7,
                value=f'=SPG("{identifier}","IQ_COMPANY_ID")')
        ws.cell(row=row, column=8,
                value=f'=SPG("{identifier}","IQ_COMPANY_NAME")')

    wb.save(XLSX_PATH)
    print(f"Created workbook: {XLSX_PATH}")


def open_refresh_and_extract():
    print("\nLaunching Excel...")
    session = launch_excel_isolated(XLSX_PATH)
    excel = session.excel
    print(f"Connected to Excel {excel.Version}")

    try:
        excel.Run("SNLXLAddin.xla!RefreshWorkbook")
        print("Triggered RefreshWorkbook")
    except Exception as e:
        print(f"Refresh warning: {e}")

    # Poll
    print("Waiting for evaluation...")
    for tick in range(30):
        time.sleep(5)
        elapsed = (tick + 1) * 5
        ws = session.workbook.Sheets(1)
        # Check a CIQ cell and an SPG cell
        ciq_val = ws.Cells(2, 3).Value
        spg_val = ws.Cells(2, 7).Value
        print(f"  [{elapsed}s] CIQ={ciq_val!r}  SPG={spg_val!r}")
        # Done when both have resolved (or errored)
        if (ciq_val is not None and str(ciq_val) != "#REFRESH" and str(ciq_val) != "#PEND"
                and spg_val is not None and str(spg_val) != "#REFRESH" and str(spg_val) != "#PEND"):
            time.sleep(10)  # let everything settle
            break
    else:
        print("  Timeout — extracting available data")

    # Extract and display
    ws = session.workbook.Sheets(1)

    def _cell(row, col):
        v = ws.Cells(row, col).Value
        if v is None:
            return "(empty)"
        s = str(v).strip()
        # Truncate long values
        return s[:30] if len(s) > 30 else s

    print("\n" + "=" * 130)
    print(f"{'Input':<22} {'Type':<14} {'CIQ() ID':<18} {'CIQRANGEA() ID':<22} {'SPG() ID':<18} {'SPG() Name':<30}")
    print("-" * 130)

    for i in range(len(TEST_IDS)):
        row = i + 2
        input_id = _cell(row, 1)
        id_type = _cell(row, 2)
        ciq_id = _cell(row, 3)
        rangea_id = _cell(row, 5)
        spg_id = _cell(row, 7)
        spg_name = _cell(row, 8)
        print(f"{input_id:<22} {id_type:<14} {ciq_id:<18} {rangea_id:<22} {spg_id:<18} {spg_name:<30}")

    print("=" * 130)

    # Also show CIQRANGEA name results to check if it returns function name
    print("\nCIQRANGEA() name column (checking for function-name-as-string bug):")
    for i in range(len(TEST_IDS)):
        row = i + 2
        input_id = _cell(row, 1)
        rangea_name = _cell(row, 6)
        print(f"  {input_id:<22} -> {rangea_name}")

    close_session(session, save=False, delete_workbook=False)


if __name__ == "__main__":
    print("=" * 60)
    print("  ID Lookup Method Comparison")
    print(f"  Testing {len(TEST_IDS)} identifiers x 3 methods")
    print("=" * 60)

    create_workbook()
    open_refresh_and_extract()
