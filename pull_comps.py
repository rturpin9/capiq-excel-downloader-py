"""
Capital IQ Comparable Companies Table — SPG Formulas.

General-purpose comp table builder supporting two modes:
  1. Lease Adjusted (default): TEV/debt INCLUDE operating leases.
     NTM EBITDA optionally adjusted for US GAAP reporters (adds TTM lease adj).
  2. Excluding Leases: TEV/debt EXCLUDE operating leases.
     EBITDA uses ex-lease-adjustment version. No NTM adjustment.

Usage:
  python pull_comps.py NYSE:HAL NYSE:SLB TSX:PD TSX:TCW
  python pull_comps.py NYSE:HAL NYSE:SLB --currency USD
  python pull_comps.py NYSE:HAL NYSE:SLB --mode excluding-leases
  python pull_comps.py NYSE:HAL NYSE:SLB --no-lease-adjust-ntm
  python pull_comps.py NYSE:HAL NYSE:SLB --groups "U.S. OFS: NYSE:HAL NYSE:SLB" "Canadian OFS: TSX:PD TSX:TCW"
"""
import argparse
import json
import os
import sys
import time

import numpy as np
import pandas as pd
from datetime import datetime
from openpyxl import Workbook

import pythoncom
from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

pythoncom.CoInitialize()

# ── Formula definitions ─────────────────────────────────────────────────────

# Lease Adjusted mode metrics
METRICS_LEASE_ADJUSTED = [
    # (label,           mnemonic,                           call_type)
    ("Company Name",    "SP_COMPANY_NAME",                  "name"),
    ("Market Cap",      "SP_MARKETCAP",                     "market"),
    ("Total Debt",      "IQ_TOTAL_DEBT",                    "bs"),
    ("Oper. Leases",    "IQ_TOTAL_OPER_LEASES",             "bs"),
    ("Cash",            "IQ_CASH_ST_INVEST",                "bs"),
    ("Net Debt",        "IQ_NET_DEBT",                      "bs"),
    ("TEV",             "IQ_TEV",                           "market"),
    ("LTM Revenue",     "IQ_TOTAL_REV",                     "ltm"),
    ("NTM Revenue",     "SP_REV_EST",                       "ntm"),
    ("LTM EBITDA",      "IQ_EBITDA_EQ_INC",                 "ltm"),
    ("Lease Adj",       "IQ_LEASE_ADJUSTMENT_EBITDA",        "ltm"),
    ("NTM EBITDA",      "SP_EBITDA_EST",                    "ntm"),
    ("GAAP",            "IQ_GAAP_BS",                       "gaap"),
    ("P/BV",            "IQ_PBV_X",                         "ltm_raw"),
]

# Excluding Leases mode metrics
METRICS_EXCLUDING_LEASES = [
    ("Company Name",    "SP_COMPANY_NAME",                  "name"),
    ("Market Cap",      "SP_MARKETCAP",                     "market"),
    ("Total Debt",      "IQ_TOTAL_DEBT_EXCL_OPER_LEASES",  "bs"),
    ("Oper. Leases",    "IQ_TOTAL_OPER_LEASES",             "bs"),
    ("Cash",            "IQ_CASH_ST_INVEST",                "bs"),
    ("TEV",             "IQ_TEV_EXCL_OPER_LEASES",          "market"),
    ("LTM Revenue",     "IQ_TOTAL_REV",                     "ltm"),
    ("NTM Revenue",     "SP_REV_EST",                       "ntm"),
    ("LTM EBITDA",      "IQ_EBITDA_EQ_INC_EXCL_OPER_LEASE_ADJ", "ltm"),
    ("NTM EBITDA",      "SP_EBITDA_EST",                    "ntm"),
    ("P/BV",            "IQ_PBV_X",                         "ltm_raw"),
]


def build_spg_formula(ticker: str, mnemonic: str, call_type: str,
                      date: str, curr_opt: str) -> str:
    t = f'"{ticker}"'
    m = f'"{mnemonic}"'
    d = f'"{date}"'
    c = f'"{curr_opt}"'

    if call_type == "name":
        return f"=SPG({t}, {m})"
    elif call_type == "market":
        return f"=SPG({t}, {m}, {d}, {c})"
    elif call_type == "bs":
        return f'=SPG({t}, {m}, "FQ0", {d}, {c})'
    elif call_type == "ltm":
        return f'=SPG({t}, {m}, "LTM", {d}, {c})'
    elif call_type == "ntm":
        return f'=SPG({t}, {m}, "NTM", {d}, {c})'
    elif call_type == "ltm_raw":
        return f'=SPG({t}, {m}, "LTM", {d})'
    elif call_type == "gaap":
        return f'=CIQ({t}, {m})'
    else:
        raise ValueError(f"Unknown call_type: {call_type}")


def build_workbook(path: str, tickers: list[str], metrics: list,
                   date: str, curr_opt: str):
    """Build workbook with primary SPG formulas + CIQRANGEA fallback.

    Layout per company (two rows):
      Row N (primary):  ticker in col A, SPG formulas using original ticker
      Row N+1 (fallback): SPG formulas using CIQRANGEA-resolved IQ ID via cell ref

    CIQRANGEA columns sit after the metrics:
      col M+2: =CIQRANGEA(ticker, "IQ_COMPANY_ID_QUICK_MATCH", 1, 1)
      col M+3: (blank — CIQRANGEA spill result = resolved IQ ID)

    The fallback row formulas reference the spill cell so SPG uses the
    resolved IQ ID. All formulas refresh in a single pass.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Comp"

    num_metrics = len(metrics)

    # Headers
    ws.cell(row=1, column=1, value="Ticker")
    for col_idx, (label, _, _) in enumerate(metrics, start=2):
        ws.cell(row=1, column=col_idx, value=label)

    # CIQRANGEA columns after metrics
    ciqrangea_col = num_metrics + 2       # formula column
    ciqrangea_spill_col = ciqrangea_col + 1  # result spills here
    ws.cell(row=1, column=ciqrangea_col, value="CIQRANGEA_Formula")
    ws.cell(row=1, column=ciqrangea_spill_col, value="Resolved_IQ_ID")

    from openpyxl.utils import get_column_letter
    spill_col_letter = get_column_letter(ciqrangea_spill_col)

    for comp_idx, ticker in enumerate(tickers):
        # Primary row: SPG formulas using original ticker
        primary_row = 2 + comp_idx * 2
        ws.cell(row=primary_row, column=1, value=ticker)
        for col_idx, (_, mnemonic, call_type) in enumerate(metrics, start=2):
            formula = build_spg_formula(ticker, mnemonic, call_type,
                                        date, curr_opt)
            ws.cell(row=primary_row, column=col_idx, value=formula)

        # CIQRANGEA lookup on primary row
        ws.cell(row=primary_row, column=ciqrangea_col,
                value=f'=CIQRANGEA("{ticker}","IQ_COMPANY_ID_QUICK_MATCH",1,1)')

        # Fallback row: SPG formulas referencing the CIQRANGEA spill cell
        # so if the original ticker fails, these use the resolved IQ ID
        fallback_row = primary_row + 1
        spill_ref = f"${spill_col_letter}${primary_row}"
        ws.cell(row=fallback_row, column=1, value=f"(fallback for {ticker})")
        for col_idx, (_, mnemonic, call_type) in enumerate(metrics, start=2):
            # Build formula that references the spill cell instead of a literal ticker
            formula = _build_spg_formula_with_ref(spill_ref, mnemonic, call_type,
                                                   date, curr_opt)
            ws.cell(row=fallback_row, column=col_idx, value=formula)

    wb.save(path)
    total_formulas = len(tickers) * num_metrics * 2  # primary + fallback
    print(f"  Created: {path} ({len(tickers)} companies x {num_metrics} metrics)")
    print(f"  {total_formulas} formulas (primary + fallback) + {len(tickers)} CIQRANGEA lookups")


def _build_spg_formula_with_ref(cell_ref: str, mnemonic: str, call_type: str,
                                 date: str, curr_opt: str) -> str:
    """Build SPG formula using a cell reference for the identifier."""
    m = f'"{mnemonic}"'
    d = f'"{date}"'
    c = f'"{curr_opt}"'

    if call_type == "name":
        return f"=SPG({cell_ref}, {m})"
    elif call_type == "market":
        return f"=SPG({cell_ref}, {m}, {d}, {c})"
    elif call_type == "bs":
        return f'=SPG({cell_ref}, {m}, "FQ0", {d}, {c})'
    elif call_type == "ltm":
        return f'=SPG({cell_ref}, {m}, "LTM", {d}, {c})'
    elif call_type == "ntm":
        return f'=SPG({cell_ref}, {m}, "NTM", {d}, {c})'
    elif call_type == "ltm_raw":
        return f'=SPG({cell_ref}, {m}, "LTM", {d})'
    elif call_type == "gaap":
        return f'=CIQ({cell_ref}, {m})'
    else:
        raise ValueError(f"Unknown call_type: {call_type}")


def safe_float(val):
    if val is None:
        return np.nan
    if isinstance(val, (int, float)):
        if val < -2000000000:
            return np.nan
        return float(val)
    if isinstance(val, str):
        upper = val.upper()
        error_tokens = ("#ERROR", "#INVALID", "#PEND", "#REFRESH", "#NAME",
                        "#OUTSIDE", "KEYERROR", "DEFUNCT", "INVALID", "NM",
                        "(INVALID")
        if any(tok in upper for tok in error_tokens):
            return np.nan
        try:
            return float(val.replace(",", ""))
        except ValueError:
            return np.nan
    return np.nan


def is_error_value(val) -> bool:
    """Check if a cell value is a CIQ/SPG error."""
    if val is None:
        return True
    if isinstance(val, (int, float)) and val < -2000000000:
        return True
    if isinstance(val, str):
        upper = val.upper()
        error_tokens = ("#ERROR", "#INVALID", "#PEND", "#REFRESH", "#NAME",
                        "#OUTSIDE", "KEYERROR", "DEFUNCT", "(INVALID")
        return any(tok in upper for tok in error_tokens)
    return False


def read_results(ws_com, num_companies: int, metrics: list) -> tuple[list[dict], list[dict]]:
    """Read results with automatic fallback to CIQRANGEA-resolved IQ IDs.

    Layout: two rows per company (primary + fallback).
      Primary row:  SPG formulas using original ticker
      Fallback row: SPG formulas using CIQRANGEA-resolved IQ ID

    For each company, uses the primary row if it resolved. If not, uses the
    fallback row (which uses the CIQRANGEA IQ ID). Reports resolution info.

    Returns:
        (results, resolution_notes) where resolution_notes is a list of
        dicts describing what happened for each ticker.
    """
    num_metrics = len(metrics)
    last_data_col = num_metrics + 1
    ciqrangea_spill_col = num_metrics + 3

    # Read the entire data region at once (all primary + fallback rows)
    total_rows = num_companies * 2
    last_row = 1 + total_rows
    data_range = ws_com.Range(
        ws_com.Cells(2, 1),
        ws_com.Cells(last_row, last_data_col)
    ).Value

    # Read CIQRANGEA spill column for all primary rows
    # (spill is only on primary rows: 2, 4, 6, ...)
    id_values = []
    for comp_idx in range(num_companies):
        primary_row = 2 + comp_idx * 2
        val = ws_com.Cells(primary_row, ciqrangea_spill_col).Value
        id_values.append(val)

    results = []
    resolution_notes = []

    for comp_idx in range(num_companies):
        primary_idx = comp_idx * 2       # index into data_range
        fallback_idx = primary_idx + 1

        primary_vals = data_range[primary_idx]
        fallback_vals = data_range[fallback_idx]
        ticker = primary_vals[0]
        resolved_id = id_values[comp_idx]

        # Check if primary SPG resolved
        primary_name = primary_vals[1]  # Company Name is first metric
        primary_ok = not is_error_value(primary_name)

        if primary_ok:
            # Primary resolved — use it
            row_data = {"_ticker": ticker}
            for met_idx, (label, _, _) in enumerate(metrics):
                row_data[label] = primary_vals[met_idx + 1]
            results.append(row_data)
            resolution_notes.append({
                "ticker": ticker, "status": "OK", "used": "primary"
            })
        else:
            # Primary failed — check fallback
            fallback_name = fallback_vals[1]
            fallback_ok = not is_error_value(fallback_name)

            if fallback_ok:
                # Fallback resolved via CIQRANGEA IQ ID
                row_data = {"_ticker": ticker}
                for met_idx, (label, _, _) in enumerate(metrics):
                    row_data[label] = fallback_vals[met_idx + 1]
                results.append(row_data)
                iq_id = str(resolved_id) if resolved_id and not is_error_value(resolved_id) else "?"
                resolution_notes.append({
                    "ticker": ticker, "status": "RESOLVED",
                    "used": "fallback", "resolved_iq_id": iq_id
                })
            else:
                # Both failed
                row_data = {"_ticker": ticker}
                for met_idx, (label, _, _) in enumerate(metrics):
                    row_data[label] = primary_vals[met_idx + 1]
                results.append(row_data)
                resolution_notes.append({
                    "ticker": ticker, "status": "FAILED", "used": "primary"
                })

    return results, resolution_notes


def safe_div(num, denom):
    if np.isnan(num) or np.isnan(denom) or denom == 0:
        return np.nan
    result = num / denom
    return result if result > 0 else np.nan


def build_table_lease_adjusted(results: list[dict], lease_adjust_ntm: bool) -> pd.DataFrame:
    rows = []
    for d in results:
        ticker = d.get("_ticker", "")
        name = d.get("Company Name", ticker)
        if not isinstance(name, str) or len(name) < 2:
            name = ticker

        mkt_cap = safe_float(d.get("Market Cap"))
        total_debt = safe_float(d.get("Total Debt"))
        oper_leases = safe_float(d.get("Oper. Leases"))
        cash = safe_float(d.get("Cash"))
        net_debt = safe_float(d.get("Net Debt"))
        tev = safe_float(d.get("TEV"))
        ltm_rev = safe_float(d.get("LTM Revenue"))
        ntm_rev = safe_float(d.get("NTM Revenue"))
        ltm_ebitda = safe_float(d.get("LTM EBITDA"))
        lease_adj = safe_float(d.get("Lease Adj"))
        ntm_ebitda_raw = safe_float(d.get("NTM EBITDA"))
        gaap = d.get("GAAP", "")
        pbv = safe_float(d.get("P/BV"))

        # Lease-adjusted NTM EBITDA
        ntm_ebitda = ntm_ebitda_raw
        adj_applied = False
        if lease_adjust_ntm and not np.isnan(ntm_ebitda_raw) and not np.isnan(lease_adj):
            if isinstance(gaap, str) and "IFRS" not in gaap.upper():
                ntm_ebitda = ntm_ebitda_raw + lease_adj
                adj_applied = True

        ebitda_margin = safe_div(ltm_ebitda, ltm_rev)
        nd_ebitda = safe_div(net_debt, ltm_ebitda)

        rows.append({
            "Company": name,
            "Ticker": ticker,
            "GAAP": gaap if isinstance(gaap, str) else "N/A",
            "Mkt Cap": mkt_cap,
            "Total Debt": total_debt,
            "Leases": oper_leases,
            "Cash": cash,
            "Net Debt": net_debt,
            "TEV": tev,
            "LTM Rev": ltm_rev,
            "NTM Rev": ntm_rev,
            "LTM EBITDA": ltm_ebitda,
            "Lease Adj": lease_adj,
            "NTM EBITDA": ntm_ebitda,
            "Lease Adj Applied": adj_applied,
            "EBITDA Margin %": ebitda_margin * 100 if not np.isnan(ebitda_margin) else np.nan,
            "EV/Rev": safe_div(tev, ltm_rev),
            "EV/LTM EBITDA": safe_div(tev, ltm_ebitda),
            "EV/NTM EBITDA": safe_div(tev, ntm_ebitda),
            "P/BV": pbv,
            "ND/EBITDA": nd_ebitda,
        })
    return pd.DataFrame(rows)


def build_table_excluding_leases(results: list[dict]) -> pd.DataFrame:
    rows = []
    for d in results:
        ticker = d.get("_ticker", "")
        name = d.get("Company Name", ticker)
        if not isinstance(name, str) or len(name) < 2:
            name = ticker

        mkt_cap = safe_float(d.get("Market Cap"))
        total_debt = safe_float(d.get("Total Debt"))  # excl oper leases
        oper_leases = safe_float(d.get("Oper. Leases"))
        cash = safe_float(d.get("Cash"))
        # Net debt calculated as debt excl leases - cash (matching Excluding Leases tab)
        net_debt = total_debt - cash if not np.isnan(total_debt) and not np.isnan(cash) else np.nan
        tev = safe_float(d.get("TEV"))  # excl oper leases
        ltm_rev = safe_float(d.get("LTM Revenue"))
        ntm_rev = safe_float(d.get("NTM Revenue"))
        ltm_ebitda = safe_float(d.get("LTM EBITDA"))  # excl lease adj
        ntm_ebitda = safe_float(d.get("NTM EBITDA"))  # straight consensus
        pbv = safe_float(d.get("P/BV"))

        ebitda_margin = safe_div(ltm_ebitda, ltm_rev)
        nd_ebitda = safe_div(net_debt, ltm_ebitda)

        rows.append({
            "Company": name,
            "Ticker": ticker,
            "Mkt Cap": mkt_cap,
            "Total Debt": total_debt,
            "Leases": oper_leases,
            "Cash": cash,
            "Net Debt": net_debt,
            "TEV": tev,
            "LTM Rev": ltm_rev,
            "NTM Rev": ntm_rev,
            "LTM EBITDA": ltm_ebitda,
            "NTM EBITDA": ntm_ebitda,
            "EBITDA Margin %": ebitda_margin * 100 if not np.isnan(ebitda_margin) else np.nan,
            "EV/Rev": safe_div(tev, ltm_rev),
            "EV/LTM EBITDA": safe_div(tev, ltm_ebitda),
            "EV/NTM EBITDA": safe_div(tev, ntm_ebitda),
            "P/BV": pbv,
            "ND/EBITDA": nd_ebitda,
        })
    return pd.DataFrame(rows)


def format_value(val, fmt_str, width):
    if pd.isna(val):
        return f"{'N/A':>{width}}"
    try:
        return f"{fmt_str.format(val):>{width}}"
    except (ValueError, TypeError):
        return f"{'N/A':>{width}}"


def print_group_stats(df, mult_cols, fmts, widths, col_order, label):
    line = f"  {label:<35}"
    for col in col_order:
        w = widths[col]
        if col in mult_cols and col in fmts:
            vals = df[col].dropna()
            if len(vals) > 0:
                fn = "mean" if "Average" in label or "Avg" in label else "median"
                s = getattr(vals, fn)()
                line += f"  {format_value(s, fmts[col], w)}"
            else:
                line += f"  {'N/A':>{w}}"
        else:
            line += f"  {'':>{w}}"
    print(line)


def format_and_print(df, mode: str, currency: str, date: str,
                     lease_adjust_ntm: bool, groups: dict | None):
    """Format and print the comp table. Returns the formatted string for output."""
    currency_upper = currency.upper()
    mode_label = "Lease Adjusted" if mode == "lease-adjusted" else "Excluding Leases"

    money = "${:,.0f}"
    mult = "{:.1f}x"
    pct = "{:.1f}%"

    if mode == "lease-adjusted":
        col_order = ["GAAP", "Mkt Cap", "Total Debt", "Leases", "Cash",
                     "Net Debt", "TEV", "LTM Rev", "NTM Rev", "LTM EBITDA",
                     "Lease Adj", "NTM EBITDA", "EBITDA Margin %", "EV/Rev",
                     "EV/LTM EBITDA", "EV/NTM EBITDA", "P/BV", "ND/EBITDA"]
        fmts = {
            "Mkt Cap": money, "Total Debt": money, "Leases": money,
            "Cash": money, "Net Debt": money, "TEV": money,
            "LTM Rev": money, "NTM Rev": money, "LTM EBITDA": money,
            "Lease Adj": money, "NTM EBITDA": money,
            "EBITDA Margin %": pct, "EV/Rev": mult,
            "EV/LTM EBITDA": mult, "EV/NTM EBITDA": mult,
            "P/BV": mult, "ND/EBITDA": mult,
        }
    else:
        col_order = ["Mkt Cap", "Total Debt", "Leases", "Cash",
                     "Net Debt", "TEV", "LTM Rev", "NTM Rev", "LTM EBITDA",
                     "NTM EBITDA", "EBITDA Margin %", "EV/Rev",
                     "EV/LTM EBITDA", "EV/NTM EBITDA", "P/BV", "ND/EBITDA"]
        fmts = {
            "Mkt Cap": money, "Total Debt": money, "Leases": money,
            "Cash": money, "Net Debt": money, "TEV": money,
            "LTM Rev": money, "NTM Rev": money, "LTM EBITDA": money,
            "NTM EBITDA": money,
            "EBITDA Margin %": pct, "EV/Rev": mult,
            "EV/LTM EBITDA": mult, "EV/NTM EBITDA": mult,
            "P/BV": mult, "ND/EBITDA": mult,
        }

    widths = {}
    for col in col_order:
        widths[col] = max(len(col), 8)
        if col == "GAAP":
            widths[col] = 9

    mult_cols = {"EBITDA Margin %", "EV/Rev", "EV/LTM EBITDA",
                 "EV/NTM EBITDA", "P/BV", "ND/EBITDA"}

    sep_len = 35 + sum(widths[c] + 2 for c in col_order) + 10
    sep = "-" * sep_len

    print(f"\n{'=' * sep_len}")
    print(f"  COMPARABLE COMPANIES ANALYSIS ({mode_label.upper()})")
    print(f"  As at {date} (Millions ${currency_upper})")
    if mode == "lease-adjusted" and lease_adjust_ntm:
        print(f"  NTM EBITDA lease-adjusted for US GAAP reporters (TTM lease adj added as proxy)")
    elif mode == "lease-adjusted":
        print(f"  NTM EBITDA: straight consensus (lease adjustment DISABLED)")
    else:
        print(f"  TEV/Debt exclude operating leases; EBITDA excludes lease adjustment")
    print(f"{'=' * sep_len}")

    def print_header():
        header = f"  {'Company':<35}"
        for col in col_order:
            w = widths[col]
            header += f"  {col:>{w}}"
        print(header)
        print(f"  {sep}")

    def print_rows(section_df):
        for _, row in section_df.iterrows():
            cname = str(row.get("Company", ""))[:33]
            line = f"  {cname:<35}"
            for col in col_order:
                w = widths[col]
                val = row.get(col, np.nan)
                if col == "GAAP":
                    s = str(val) if isinstance(val, str) else "N/A"
                    line += f"  {s:>{w}}"
                elif col in fmts:
                    line += f"  {format_value(val, fmts[col], w)}"
                else:
                    line += f"  {str(val):>{w}}"
            if mode == "lease-adjusted" and row.get("Lease Adj Applied", False):
                line += "  *adj"
            print(line)

    if groups:
        for group_name, group_tickers in groups.items():
            group_df = df[df["Ticker"].isin(group_tickers)]
            if group_df.empty:
                continue

            print(f"\n  {group_name}")
            print(f"  {sep}")
            print_header()
            print_rows(group_df)
            print(f"  {sep}")

            # Group stats
            print_group_stats(group_df, mult_cols, fmts, widths, col_order,
                              f"Average {group_name}")
            print_group_stats(group_df, mult_cols, fmts, widths, col_order,
                              f"Median {group_name}")
            # Swap fn for median
    else:
        print_header()
        print_rows(df)
        print(f"  {sep}")

    # Overall
    print(f"\n  Overall Summary")
    print(f"  {sep}")
    for stat_label, stat_fn in [("Overall Average", "mean"), ("Overall Median", "median")]:
        line = f"  {stat_label:<35}"
        for col in col_order:
            w = widths[col]
            if col in mult_cols and col in fmts:
                vals = df[col].dropna()
                if len(vals) > 0:
                    s = getattr(vals, stat_fn)()
                    line += f"  {format_value(s, fmts[col], w)}"
                else:
                    line += f"  {'N/A':>{w}}"
            else:
                line += f"  {'':>{w}}"
        print(line)

    print(f"{'=' * sep_len}")

    if mode == "lease-adjusted" and lease_adjust_ntm:
        print("  * adj = NTM EBITDA lease-adjusted (TTM lease expense added for US GAAP reporters)")
        print("  IFRS 16 reporters already have leases below EBITDA; no adjustment needed.")


def print_group_stats(df, mult_cols, fmts, widths, col_order, label):
    fn = "mean" if "Average" in label or "Avg" in label else "median"
    line = f"  {label:<35}"
    for col in col_order:
        w = widths[col]
        if col in mult_cols and col in fmts:
            vals = df[col].dropna()
            if len(vals) > 0:
                s = getattr(vals, fn)()
                line += f"  {format_value(s, fmts[col], w)}"
            else:
                line += f"  {'N/A':>{w}}"
        else:
            line += f"  {'':>{w}}"
    print(line)


# ── JSON output for skill consumption ───────────────────────────────────────

def results_to_json(df, mode, currency, date, lease_adjust_ntm):
    """Return structured JSON for programmatic consumption."""
    output = {
        "mode": mode,
        "currency": currency,
        "date": date,
        "lease_adjust_ntm": lease_adjust_ntm if mode == "lease-adjusted" else None,
        "companies": [],
    }
    for _, row in df.iterrows():
        company = {}
        for col in df.columns:
            val = row[col]
            if isinstance(val, float) and np.isnan(val):
                company[col] = None
            elif isinstance(val, (np.floating, np.integer)):
                company[col] = float(val)
            else:
                company[col] = val
        output["companies"].append(company)
    return output


# ── Main ────────────────────────────────────────────────────────────────────

def parse_groups(group_args: list[str] | None, all_tickers: list[str]) -> dict | None:
    """Parse --groups 'Label: T1 T2' 'Label2: T3 T4' into dict."""
    if not group_args:
        return None

    groups = {}
    for g in group_args:
        if ":" not in g:
            print(f"  Warning: ignoring malformed group '{g}' (expected 'Label: T1 T2 ...')")
            continue
        label, tickers_str = g.split(":", 1)
        tickers = tickers_str.strip().split()
        groups[label.strip()] = tickers
    return groups if groups else None


def main():
    parser = argparse.ArgumentParser(
        description="Pull comparable company data via S&P Capital IQ SPG formulas.")
    parser.add_argument("tickers", nargs="+",
                        help="Company tickers (e.g. NYSE:HAL TSX:PD DSGX)")
    parser.add_argument("--currency", default="CAD",
                        help="Currency for output (default: CAD)")
    parser.add_argument("--mode", choices=["lease-adjusted", "excluding-leases"],
                        default="lease-adjusted",
                        help="Comp table mode (default: lease-adjusted)")
    parser.add_argument("--no-lease-adjust-ntm", action="store_true",
                        help="Disable NTM EBITDA lease adjustment for US GAAP reporters")
    parser.add_argument("--groups", nargs="*",
                        help='Group tickers: "U.S.: NYSE:HAL NYSE:SLB" "Canada: TSX:PD TSX:TCW"')
    parser.add_argument("--date", default=None,
                        help="As-of date in M/D/YYYY format (default: today)")
    parser.add_argument("--json-output", default=None,
                        help="Write structured JSON to this path")
    parser.add_argument("--csv-output", default=None,
                        help="Write CSV to this path (default: comps_output.csv)")
    parser.add_argument("--max-wait", type=int, default=180,
                        help="Max seconds to wait for formula refresh (default: 180)")

    args = parser.parse_args()

    # Resolve date
    if args.date:
        date_str = args.date
    else:
        now = datetime.now()
        date_str = f"{now.month}/{now.day}/{now.year}"

    curr_opt = f"Options: Curr={args.currency.upper()}"
    lease_adjust_ntm = not args.no_lease_adjust_ntm
    mode = args.mode
    tickers = args.tickers

    # Select metrics for mode
    if mode == "lease-adjusted":
        metrics = METRICS_LEASE_ADJUSTED
    else:
        metrics = METRICS_EXCLUDING_LEASES

    groups = parse_groups(args.groups, tickers)

    xlsx_path = os.path.abspath("_comps_temp.xlsx")
    csv_path = args.csv_output or os.path.abspath("comps_output.csv")

    print(f"Comparable Companies Analysis")
    print(f"  Mode: {mode}")
    print(f"  Currency: {args.currency.upper()}")
    print(f"  Date: {date_str}")
    print(f"  Tickers: {' '.join(tickers)}")
    if mode == "lease-adjusted":
        print(f"  NTM lease adjustment: {'ON' if lease_adjust_ntm else 'OFF'}")
    print()

    # Build workbook
    print("Building workbook...")
    build_workbook(xlsx_path, tickers, metrics, date_str, curr_opt)

    # Launch Excel
    print("\nLaunching Excel (isolated instance)...")
    session = launch_excel_isolated(xlsx_path)
    excel = session.excel
    print(f"  Connected to Excel {excel.Version}")
    ws_com = session.workbook.Sheets(1)

    # Trigger refresh
    print("Triggering RefreshSheet...")
    try:
        excel.Run("SNLXLAddin.xla!RefreshSheet")
        print("  RefreshSheet executed")
    except Exception as e:
        print(f"  RefreshSheet warning: {e}")

    # Poll for completion (two rows per company: primary + fallback)
    num_formulas = len(tickers) * len(metrics) * 2
    print(f"Waiting for formulas (max {args.max_wait}s)...")
    start = time.monotonic()

    while True:
        elapsed = time.monotonic() - start
        if elapsed > args.max_wait:
            print(f"  Timeout after {args.max_wait}s -- proceeding with available data")
            break

        pending = False
        resolved = 0
        for comp_idx in range(len(tickers)):
            for row_offset in (0, 1):  # primary + fallback
                row = 2 + comp_idx * 2 + row_offset
                for met_idx in range(len(metrics)):
                    col = met_idx + 2
                    val = ws_com.Cells(row, col).Value
                    if isinstance(val, str) and val.upper() in ('#PEND', '#REFRESH'):
                        pending = True
                    elif val is not None and not (isinstance(val, (int, float)) and val < -2000000000):
                        resolved += 1

        if not pending and resolved > num_formulas * 0.5:
            print(f"  Data ready! {resolved}/{num_formulas} cells resolved in {elapsed:.0f}s")
            break

        if int(elapsed) % 15 < 5:
            print(f"  {elapsed:.0f}s: {resolved}/{num_formulas} resolved, pending={pending}")

        time.sleep(5)

    # Read results
    print("\nReading data...")
    results, resolution_notes = read_results(ws_com, len(tickers), metrics)

    # Close Excel
    close_session(session, save=False, delete_workbook=True)

    # Report identifier resolution
    resolved_via_fallback = [n for n in resolution_notes if n["status"] == "RESOLVED"]
    failed = [n for n in resolution_notes if n["status"] == "FAILED"]

    if resolved_via_fallback or failed:
        print("\n" + "=" * 70)
        print("  IDENTIFIER RESOLUTION REPORT")
        print("=" * 70)
        for note in resolved_via_fallback:
            iq_id = note.get("resolved_iq_id", "?")
            print(f"  {note['ticker']}: resolved via CIQRANGEA -> {iq_id} (data OK)")
        for note in failed:
            print(f"  {note['ticker']}: FAILED — not recognized by SPG or CIQRANGEA")
            print(f"    -> Check spelling, or try EXCHANGE:TICKER format (e.g., NYSE:HAL, TSX:PD)")
        print("=" * 70 + "\n")

    # Build table
    if mode == "lease-adjusted":
        df = build_table_lease_adjusted(results, lease_adjust_ntm)
    else:
        df = build_table_excluding_leases(results)

    # Print formatted output
    format_and_print(df, mode, args.currency, date_str, lease_adjust_ntm, groups)

    # Save CSV
    df.to_csv(csv_path, index=False)
    print(f"\nCSV saved to: {csv_path}")

    # Save JSON if requested
    if args.json_output:
        data = results_to_json(df, mode, args.currency, date_str, lease_adjust_ntm)
        with open(args.json_output, "w") as f:
            json.dump(data, f, indent=2)
        print(f"JSON saved to: {args.json_output}")

    return df


if __name__ == "__main__":
    main()
