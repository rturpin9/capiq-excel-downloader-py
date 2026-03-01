"""Lightweight identifier resolution via CIQRANGEA + CIQ company name lookup.

Builds a small workbook with just ID resolution formulas, launches Excel,
reads results. Much faster than a full comp pull (~10s vs ~40s).
"""
from __future__ import annotations

import logging
import os
import time

from openpyxl import Workbook

log = logging.getLogger("capiq_mcp")


def build_id_workbook(path: str, identifiers: list[str]) -> None:
    """Build a workbook with CIQRANGEA ID lookup + CIQ company name per identifier.

    Layout:
      Col A: identifier (raw input)
      Col B: =CIQRANGEA(id, "IQ_COMPANY_ID_QUICK_MATCH", 1, 1)  (formula)
      Col C: (blank — CIQRANGEA spill result = resolved IQ ID)
      Col D: =CIQ(spill_ref, "IQ_COMPANY_NAME")  (name from resolved ID)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Lookup"

    ws.cell(row=1, column=1, value="Input")
    ws.cell(row=1, column=2, value="CIQRANGEA_Formula")
    ws.cell(row=1, column=3, value="Resolved_IQ_ID")
    ws.cell(row=1, column=4, value="Company_Name")

    for idx, ident in enumerate(identifiers):
        row = idx + 2
        # CIQRANGEA doesn't handle EXCHANGE:TICKER format — strip prefix
        lookup_id = ident.split(":", 1)[1] if ":" in ident else ident
        ws.cell(row=row, column=1, value=ident)
        ws.cell(row=row, column=2,
                value=f'=CIQRANGEA("{lookup_id}","IQ_COMPANY_ID_QUICK_MATCH",1,1)')
        # Col C is blank — CIQRANGEA result spills here
        # Company name from resolved IQ ID (references spill cell)
        ws.cell(row=row, column=4,
                value=f'=CIQ($C${row},"IQ_COMPANY_NAME")')

    wb.save(path)
    log.info("Created ID lookup workbook: %s (%d identifiers)", path, len(identifiers))


def run_id_lookup(identifiers: list[str], max_wait: int = 60) -> dict:
    """Resolve identifiers to CIQ IQ IDs and company names.

    Parameters
    ----------
    identifiers : list[str]
        Tickers, company names, CUSIPs, ISINs, etc.
    max_wait : int
        Max seconds to wait for formula refresh.

    Returns
    -------
    dict
        {"results": [{"input": ..., "iq_id": ..., "company_name": ..., "status": ...}]}
    """
    import pythoncom
    pythoncom.CoInitialize()

    from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

    xlsx_path = os.path.abspath("_id_lookup_mcp_temp.xlsx")
    build_id_workbook(xlsx_path, identifiers)

    session = launch_excel_isolated(xlsx_path)
    try:
        excel = session.excel
        log.info("Connected to Excel %s for ID lookup", excel.Version)
        ws_com = session.workbook.Sheets(1)

        # Trigger refresh
        import pythoncom
        log.info("Calling RefreshSheet for ID lookup...")
        try:
            pythoncom.PumpWaitingMessages()
            excel.Run("SNLXLAddin.xla!RefreshSheet")
            log.info("RefreshSheet executed for ID lookup")
        except Exception as e:
            log.warning("RefreshSheet warning: %s", e)
        pythoncom.PumpWaitingMessages()

        # Poll for completion — batch read CIQRANGEA spill cells.
        # Exit when ALL cells resolve, OR when the resolved count
        # stabilizes (stops increasing for 3 consecutive polls),
        # meaning remaining cells are stuck on #PEND/#ERROR.
        n = len(identifiers)
        start = time.monotonic()
        prev_resolved = -1
        stable_count = 0

        while True:
            elapsed = time.monotonic() - start
            if elapsed > max_wait:
                log.warning("ID lookup timeout after %ds", max_wait)
                break

            # Single batch read for spill column
            if n == 1:
                spill_vals = ((ws_com.Cells(2, 3).Value,),)
            else:
                spill_vals = ws_com.Range(
                    ws_com.Cells(2, 3), ws_com.Cells(1 + n, 3)
                ).Value

            resolved = 0
            pending = 0
            for row_tuple in spill_vals:
                val = row_tuple[0]
                if val is None:
                    pending += 1
                elif isinstance(val, str) and val.upper() in ('#PEND', '#REFRESH'):
                    pending += 1
                else:
                    resolved += 1

            log.debug("ID lookup %ds: %d/%d resolved, %d pending",
                      elapsed, resolved, n, pending)

            if pending == 0:
                log.info("ID lookup fully resolved in %.0fs (%d/%d)", elapsed, resolved, n)
                break

            # If resolved count hasn't changed for 3 polls (~6s), remaining
            # cells are stuck — proceed with what we have.
            if resolved == prev_resolved and resolved > 0:
                stable_count += 1
                if stable_count >= 3:
                    log.info("ID lookup settled in %.0fs (%d/%d resolved, %d stuck)",
                             elapsed, resolved, n, pending)
                    break
            else:
                stable_count = 0
            prev_resolved = resolved

            # Pump COM message queue to prevent STA deadlocks
            pythoncom.PumpWaitingMessages()
            time.sleep(2)

        # Batch read results (columns C-D)
        if n == 1:
            result_vals = ((ws_com.Cells(2, 3).Value, ws_com.Cells(2, 4).Value),)
        else:
            result_vals = ws_com.Range(
                ws_com.Cells(2, 3), ws_com.Cells(1 + n, 4)
            ).Value

        results = []
        for idx, ident in enumerate(identifiers):
            iq_id = result_vals[idx][0]
            company_name = result_vals[idx][1]

            # Classify result
            if iq_id is None or (isinstance(iq_id, str) and
                                  any(tok in iq_id.upper() for tok in
                                      ("#ERROR", "#INVALID", "#NAME", "KEYERROR"))):
                results.append({
                    "input": ident,
                    "iq_id": None,
                    "company_name": None,
                    "status": "FAILED",
                })
            else:
                name = None
                if company_name and isinstance(company_name, str):
                    upper = company_name.upper()
                    if not any(tok in upper for tok in ("#ERROR", "#INVALID", "#NAME")):
                        name = company_name
                results.append({
                    "input": ident,
                    "iq_id": str(iq_id),
                    "company_name": name,
                    "status": "OK",
                })
    finally:
        close_session(session, save=False, delete_workbook=True)

    log.info("ID lookup complete: %d/%d resolved",
             sum(1 for r in results if r["status"] == "OK"), len(results))
    return {"results": results}
