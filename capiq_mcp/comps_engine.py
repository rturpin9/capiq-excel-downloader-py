"""Core comp table engine — extracted from pull_comps.py.

All business logic for building, refreshing, and reading CIQ comp tables.
No print() — uses logging throughout. Returns structured dicts for MCP consumption.
"""
from __future__ import annotations

import logging
import os
import time
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

log = logging.getLogger("capiq_mcp")

# ── Metric definitions ─────────────────────────────────────────────────────

METRICS_LEASE_ADJUSTED = [
    # (label,           mnemonic,                           call_type)
    ("Company Name",    "SP_COMPANY_NAME",                  "name"),
    ("Market Cap",      "SP_MARKETCAP",                     "market"),
    ("Total Debt",      "IQ_TOTAL_DEBT",                    "bs"),
    ("Oper. Leases",    "IQ_TOTAL_OPER_LEASES",             "bs"),
    ("Cash",            "IQ_CASH_ST_INVEST",                "bs"),
    ("Net Debt",        "IQ_NET_DEBT",                      "bs"),
    ("TEV",             "IQ_TEV",                           "market"),
    ("LTM Revenue",     "IQ_TOTAL_REV",                     "ltm"),
    ("NTM Revenue",     "SP_REV_EST",                       "ntm"),
    ("LTM EBITDA",      "IQ_EBITDA_EQ_INC",                 "ltm"),
    ("Lease Adj",       "IQ_LEASE_ADJUSTMENT_EBITDA",        "ltm"),
    ("NTM EBITDA",      "SP_EBITDA_EST",                    "ntm"),
    ("GAAP",            "IQ_GAAP_BS",                       "gaap"),
    ("P/BV",            "IQ_PBV_X",                         "ltm_raw"),
]

METRICS_EXCLUDING_LEASES = [
    ("Company Name",    "SP_COMPANY_NAME",                  "name"),
    ("Market Cap",      "SP_MARKETCAP",                     "market"),
    ("Total Debt",      "IQ_TOTAL_DEBT_EXCL_OPER_LEASES",  "bs"),
    ("Oper. Leases",    "IQ_TOTAL_OPER_LEASES",             "bs"),
    ("Cash",            "IQ_CASH_ST_INVEST",                "bs"),
    ("TEV",             "IQ_TEV_EXCL_OPER_LEASES",          "market"),
    ("LTM Revenue",     "IQ_TOTAL_REV",                     "ltm"),
    ("NTM Revenue",     "SP_REV_EST",                       "ntm"),
    ("LTM EBITDA",      "IQ_EBITDA_EQ_INC_EXCL_OPER_LEASE_ADJ", "ltm"),
    ("NTM EBITDA",      "SP_EBITDA_EST",                    "ntm"),
    ("P/BV",            "IQ_PBV_X",                         "ltm_raw"),
]


# ── Formula builders ───────────────────────────────────────────────────────

def build_spg_formula(ticker: str, mnemonic: str, call_type: str,
                      date: str, curr_opt: str) -> str:
    """Build an SPG formula string for a given metric type."""
    t = f'"{ticker}"'
    m = f'"{mnemonic}"'
    d = f'"{date}"'
    c = f'"{curr_opt}"'

    if call_type == "name":
        return f"=SPG({t}, {m})"
    elif call_type == "market":
        return f"=SPG({t}, {m}, {d}, {c})"
    elif call_type == "bs":
        return f'=SPG({t}, {m}, "FQ0", {d}, {c})'
    elif call_type == "ltm":
        return f'=SPG({t}, {m}, "LTM", {d}, {c})'
    elif call_type == "ntm":
        return f'=SPG({t}, {m}, "NTM", {d}, {c})'
    elif call_type == "ltm_raw":
        return f'=SPG({t}, {m}, "LTM", {d})'
    elif call_type == "gaap":
        return f'=CIQ({t}, {m})'
    else:
        raise ValueError(f"Unknown call_type: {call_type}")


def _build_spg_formula_with_ref(cell_ref: str, mnemonic: str, call_type: str,
                                 date: str, curr_opt: str) -> str:
    """Build SPG formula using a cell reference for the identifier."""
    m = f'"{mnemonic}"'
    d = f'"{date}"'
    c = f'"{curr_opt}"'

    if call_type == "name":
        return f"=SPG({cell_ref}, {m})"
    elif call_type == "market":
        return f"=SPG({cell_ref}, {m}, {d}, {c})"
    elif call_type == "bs":
        return f'=SPG({cell_ref}, {m}, "FQ0", {d}, {c})'
    elif call_type == "ltm":
        return f'=SPG({cell_ref}, {m}, "LTM", {d}, {c})'
    elif call_type == "ntm":
        return f'=SPG({cell_ref}, {m}, "NTM", {d}, {c})'
    elif call_type == "ltm_raw":
        return f'=SPG({cell_ref}, {m}, "LTM", {d})'
    elif call_type == "gaap":
        return f'=CIQ({cell_ref}, {m})'
    else:
        raise ValueError(f"Unknown call_type: {call_type}")


# ── Workbook creation ──────────────────────────────────────────────────────

def build_workbook(path: str, tickers: list[str], metrics: list,
                   date: str, curr_opt: str) -> None:
    """Build workbook with primary SPG formulas + CIQRANGEA fallback.

    Layout per company (two rows):
      Row N (primary):  ticker in col A, SPG formulas using original ticker
      Row N+1 (fallback): SPG formulas using CIQRANGEA-resolved IQ ID via cell ref

    CIQRANGEA columns sit after the metrics:
      col M+2: =CIQRANGEA(ticker, "IQ_COMPANY_ID_QUICK_MATCH", 1, 1)
      col M+3: (blank — CIQRANGEA spill result = resolved IQ ID)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Comp"

    num_metrics = len(metrics)

    # Headers
    ws.cell(row=1, column=1, value="Ticker")
    for col_idx, (label, _, _) in enumerate(metrics, start=2):
        ws.cell(row=1, column=col_idx, value=label)

    # CIQRANGEA columns after metrics
    ciqrangea_col = num_metrics + 2
    ciqrangea_spill_col = ciqrangea_col + 1
    ws.cell(row=1, column=ciqrangea_col, value="CIQRANGEA_Formula")
    ws.cell(row=1, column=ciqrangea_spill_col, value="Resolved_IQ_ID")

    spill_col_letter = get_column_letter(ciqrangea_spill_col)

    for comp_idx, ticker in enumerate(tickers):
        # Primary row: SPG formulas using original ticker
        primary_row = 2 + comp_idx * 2
        ws.cell(row=primary_row, column=1, value=ticker)
        for col_idx, (_, mnemonic, call_type) in enumerate(metrics, start=2):
            formula = build_spg_formula(ticker, mnemonic, call_type,
                                        date, curr_opt)
            ws.cell(row=primary_row, column=col_idx, value=formula)

        # CIQRANGEA lookup on primary row
        # CIQRANGEA doesn't handle EXCHANGE:TICKER format — strip prefix
        lookup_id = ticker.split(":", 1)[1] if ":" in ticker else ticker
        ws.cell(row=primary_row, column=ciqrangea_col,
                value=f'=CIQRANGEA("{lookup_id}","IQ_COMPANY_ID_QUICK_MATCH",1,1)')

        # Fallback row: SPG formulas referencing the CIQRANGEA spill cell
        fallback_row = primary_row + 1
        spill_ref = f"${spill_col_letter}${primary_row}"
        ws.cell(row=fallback_row, column=1, value=f"(fallback for {ticker})")
        for col_idx, (_, mnemonic, call_type) in enumerate(metrics, start=2):
            formula = _build_spg_formula_with_ref(spill_ref, mnemonic, call_type,
                                                   date, curr_opt)
            ws.cell(row=fallback_row, column=col_idx, value=formula)

    wb.save(path)
    total_formulas = len(tickers) * num_metrics * 2
    log.info("Created %s (%d companies x %d metrics, %d formulas + %d CIQRANGEA lookups)",
             path, len(tickers), num_metrics, total_formulas, len(tickers))


# ── Value helpers ──────────────────────────────────────────────────────────

def safe_float(val) -> float:
    """Convert a CIQ cell value to float, returning NaN for errors."""
    if val is None:
        return np.nan
    if isinstance(val, (int, float)):
        if val < -2000000000:
            return np.nan
        return float(val)
    if isinstance(val, str):
        upper = val.upper()
        error_tokens = ("#ERROR", "#INVALID", "#PEND", "#REFRESH", "#NAME",
                        "#OUTSIDE", "KEYERROR", "DEFUNCT", "INVALID", "NM",
                        "(INVALID")
        if any(tok in upper for tok in error_tokens):
            return np.nan
        try:
            return float(val.replace(",", ""))
        except ValueError:
            return np.nan
    return np.nan


def is_error_value(val) -> bool:
    """Check if a cell value is a CIQ/SPG error."""
    if val is None:
        return True
    if isinstance(val, (int, float)) and val < -2000000000:
        return True
    if isinstance(val, str):
        upper = val.upper()
        error_tokens = ("#ERROR", "#INVALID", "#PEND", "#REFRESH", "#NAME",
                        "#OUTSIDE", "KEYERROR", "DEFUNCT", "(INVALID")
        return any(tok in upper for tok in error_tokens)
    return False


def safe_div(num: float, denom: float) -> float:
    """Safe division returning NaN for invalid or non-positive results."""
    if np.isnan(num) or np.isnan(denom) or denom == 0:
        return np.nan
    result = num / denom
    return result if result > 0 else np.nan


# ── Result reading ─────────────────────────────────────────────────────────

def read_results(ws_com, num_companies: int, metrics: list) -> tuple[list[dict], list[dict]]:
    """Read results with automatic fallback to CIQRANGEA-resolved IQ IDs.

    Returns (results, resolution_notes).
    """
    num_metrics = len(metrics)
    last_data_col = num_metrics + 1
    ciqrangea_spill_col = num_metrics + 3

    # Read the entire data region at once (batch COM read)
    total_rows = num_companies * 2
    last_row = 1 + total_rows
    data_range = ws_com.Range(
        ws_com.Cells(2, 1),
        ws_com.Cells(last_row, last_data_col)
    ).Value

    # Read CIQRANGEA spill column for all primary rows
    id_values = []
    for comp_idx in range(num_companies):
        primary_row = 2 + comp_idx * 2
        val = ws_com.Cells(primary_row, ciqrangea_spill_col).Value
        id_values.append(val)

    results = []
    resolution_notes = []

    for comp_idx in range(num_companies):
        primary_idx = comp_idx * 2
        fallback_idx = primary_idx + 1

        primary_vals = data_range[primary_idx]
        fallback_vals = data_range[fallback_idx]
        ticker = primary_vals[0]
        resolved_id = id_values[comp_idx]

        # Check if primary SPG resolved
        primary_name = primary_vals[1]  # Company Name is first metric
        primary_ok = not is_error_value(primary_name)

        if primary_ok:
            row_data = {"_ticker": ticker}
            for met_idx, (label, _, _) in enumerate(metrics):
                row_data[label] = primary_vals[met_idx + 1]
            results.append(row_data)
            resolution_notes.append({
                "ticker": ticker, "status": "OK", "used": "primary"
            })
        else:
            fallback_name = fallback_vals[1]
            fallback_ok = not is_error_value(fallback_name)

            if fallback_ok:
                row_data = {"_ticker": ticker}
                for met_idx, (label, _, _) in enumerate(metrics):
                    row_data[label] = fallback_vals[met_idx + 1]
                results.append(row_data)
                iq_id = str(resolved_id) if resolved_id and not is_error_value(resolved_id) else "?"
                resolution_notes.append({
                    "ticker": ticker, "status": "RESOLVED",
                    "used": "fallback", "resolved_iq_id": iq_id
                })
            else:
                row_data = {"_ticker": ticker}
                for met_idx, (label, _, _) in enumerate(metrics):
                    row_data[label] = primary_vals[met_idx + 1]
                results.append(row_data)
                resolution_notes.append({
                    "ticker": ticker, "status": "FAILED", "used": "primary"
                })

    return results, resolution_notes


# ── Table builders ─────────────────────────────────────────────────────────

def build_table_lease_adjusted(results: list[dict], lease_adjust_ntm: bool) -> pd.DataFrame:
    """Build the lease-adjusted comp table DataFrame."""
    rows = []
    for d in results:
        ticker = d.get("_ticker", "")
        name = d.get("Company Name", ticker)
        if not isinstance(name, str) or len(name) < 2:
            name = ticker

        mkt_cap = safe_float(d.get("Market Cap"))
        total_debt = safe_float(d.get("Total Debt"))
        oper_leases = safe_float(d.get("Oper. Leases"))
        cash = safe_float(d.get("Cash"))
        net_debt = safe_float(d.get("Net Debt"))
        tev = safe_float(d.get("TEV"))
        ltm_rev = safe_float(d.get("LTM Revenue"))
        ntm_rev = safe_float(d.get("NTM Revenue"))
        ltm_ebitda = safe_float(d.get("LTM EBITDA"))
        lease_adj = safe_float(d.get("Lease Adj"))
        ntm_ebitda_raw = safe_float(d.get("NTM EBITDA"))
        gaap = d.get("GAAP", "")
        pbv = safe_float(d.get("P/BV"))

        # Lease-adjusted NTM EBITDA
        ntm_ebitda = ntm_ebitda_raw
        adj_applied = False
        if lease_adjust_ntm and not np.isnan(ntm_ebitda_raw) and not np.isnan(lease_adj):
            if isinstance(gaap, str) and "IFRS" not in gaap.upper():
                ntm_ebitda = ntm_ebitda_raw + lease_adj
                adj_applied = True

        ebitda_margin = safe_div(ltm_ebitda, ltm_rev)
        nd_ebitda = safe_div(net_debt, ltm_ebitda)

        rows.append({
            "Company": name,
            "Ticker": ticker,
            "GAAP": gaap if isinstance(gaap, str) else "N/A",
            "Mkt Cap": mkt_cap,
            "Total Debt": total_debt,
            "Leases": oper_leases,
            "Cash": cash,
            "Net Debt": net_debt,
            "TEV": tev,
            "LTM Rev": ltm_rev,
            "NTM Rev": ntm_rev,
            "LTM EBITDA": ltm_ebitda,
            "Lease Adj": lease_adj,
            "NTM EBITDA": ntm_ebitda,
            "Lease Adj Applied": adj_applied,
            "EBITDA Margin %": ebitda_margin * 100 if not np.isnan(ebitda_margin) else np.nan,
            "EV/Rev": safe_div(tev, ltm_rev),
            "EV/LTM EBITDA": safe_div(tev, ltm_ebitda),
            "EV/NTM EBITDA": safe_div(tev, ntm_ebitda),
            "P/BV": pbv,
            "ND/EBITDA": nd_ebitda,
        })
    return pd.DataFrame(rows)


def build_table_excluding_leases(results: list[dict]) -> pd.DataFrame:
    """Build the excluding-leases comp table DataFrame."""
    rows = []
    for d in results:
        ticker = d.get("_ticker", "")
        name = d.get("Company Name", ticker)
        if not isinstance(name, str) or len(name) < 2:
            name = ticker

        mkt_cap = safe_float(d.get("Market Cap"))
        total_debt = safe_float(d.get("Total Debt"))
        oper_leases = safe_float(d.get("Oper. Leases"))
        cash = safe_float(d.get("Cash"))
        net_debt = total_debt - cash if not np.isnan(total_debt) and not np.isnan(cash) else np.nan
        tev = safe_float(d.get("TEV"))
        ltm_rev = safe_float(d.get("LTM Revenue"))
        ntm_rev = safe_float(d.get("NTM Revenue"))
        ltm_ebitda = safe_float(d.get("LTM EBITDA"))
        ntm_ebitda = safe_float(d.get("NTM EBITDA"))
        pbv = safe_float(d.get("P/BV"))

        ebitda_margin = safe_div(ltm_ebitda, ltm_rev)
        nd_ebitda = safe_div(net_debt, ltm_ebitda)

        rows.append({
            "Company": name,
            "Ticker": ticker,
            "Mkt Cap": mkt_cap,
            "Total Debt": total_debt,
            "Leases": oper_leases,
            "Cash": cash,
            "Net Debt": net_debt,
            "TEV": tev,
            "LTM Rev": ltm_rev,
            "NTM Rev": ntm_rev,
            "LTM EBITDA": ltm_ebitda,
            "NTM EBITDA": ntm_ebitda,
            "EBITDA Margin %": ebitda_margin * 100 if not np.isnan(ebitda_margin) else np.nan,
            "EV/Rev": safe_div(tev, ltm_rev),
            "EV/LTM EBITDA": safe_div(tev, ltm_ebitda),
            "EV/NTM EBITDA": safe_div(tev, ntm_ebitda),
            "P/BV": pbv,
            "ND/EBITDA": nd_ebitda,
        })
    return pd.DataFrame(rows)


# ── JSON serialization ─────────────────────────────────────────────────────

def dataframe_to_json(df: pd.DataFrame, mode: str, currency: str,
                      date: str, lease_adjust_ntm: bool) -> dict:
    """Convert a comp table DataFrame to a structured dict for MCP output."""
    output = {
        "mode": mode,
        "currency": currency,
        "date": date,
        "lease_adjust_ntm": lease_adjust_ntm if mode == "lease-adjusted" else None,
        "companies": [],
    }
    for _, row in df.iterrows():
        company = {}
        for col in df.columns:
            val = row[col]
            if isinstance(val, float) and np.isnan(val):
                company[col] = None
            elif isinstance(val, (np.floating, np.integer)):
                company[col] = float(val)
            else:
                company[col] = val
        output["companies"].append(company)
    return output


# ── Polling ────────────────────────────────────────────────────────────────

def poll_for_completion(ws_com, num_companies: int, num_metrics: int,
                        max_wait: float) -> None:
    """Poll formula cells until data is ready or timeout.

    Uses a single batch COM read per cycle instead of cell-by-cell access.
    """
    num_formulas = num_companies * num_metrics * 2  # primary + fallback
    log.info("Waiting for %d formulas (max %ds)...", num_formulas, max_wait)
    start = time.monotonic()

    total_rows = num_companies * 2
    first_row = 2
    last_row = 1 + total_rows
    first_col = 2
    last_col = num_metrics + 1

    while True:
        elapsed = time.monotonic() - start
        if elapsed > max_wait:
            log.warning("Timeout after %ds — proceeding with available data", max_wait)
            break

        # Single batch COM read for all formula cells
        data = ws_com.Range(
            ws_com.Cells(first_row, first_col),
            ws_com.Cells(last_row, last_col),
        ).Value

        pending = False
        resolved = 0
        for row_data in data:
            for val in row_data:
                if isinstance(val, str) and val.upper() in ('#PEND', '#REFRESH'):
                    pending = True
                elif val is not None and not (isinstance(val, (int, float)) and val < -2000000000):
                    resolved += 1

        if not pending and resolved > num_formulas * 0.5:
            log.info("Data ready: %d/%d cells resolved in %.0fs", resolved, num_formulas, elapsed)
            break

        if int(elapsed) % 15 < 5:
            log.debug("%ds: %d/%d resolved, pending=%s", elapsed, resolved, num_formulas, pending)

        time.sleep(5)


# ── Main pipeline ──────────────────────────────────────────────────────────

def run_comps(
    tickers: list[str],
    currency: str = "CAD",
    mode: str = "lease-adjusted",
    lease_adjust_ntm: bool = True,
    date: str | None = None,
    max_wait: int = 180,
) -> dict:
    """Full pipeline: build workbook -> launch Excel -> refresh -> read -> return JSON.

    Parameters
    ----------
    tickers : list[str]
        Company identifiers (EXCHANGE:TICKER format preferred).
    currency : str
        Output currency (default CAD).
    mode : str
        "lease-adjusted" or "excluding-leases".
    lease_adjust_ntm : bool
        Whether to add TTM lease adj to NTM EBITDA for US GAAP reporters.
    date : str or None
        As-of date in M/D/YYYY format. Defaults to today.
    max_wait : int
        Max seconds to wait for formula refresh.

    Returns
    -------
    dict
        Structured result with companies array and resolution_notes.
    """
    import pythoncom
    pythoncom.CoInitialize()

    from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

    # Resolve date
    if date is None:
        now = datetime.now()
        date = f"{now.month}/{now.day}/{now.year}"

    curr_opt = f"Options: Curr={currency.upper()}"

    # Select metrics
    metrics = METRICS_LEASE_ADJUSTED if mode == "lease-adjusted" else METRICS_EXCLUDING_LEASES

    xlsx_path = os.path.abspath("_comps_mcp_temp.xlsx")

    log.info("Comp pull: %d tickers, mode=%s, currency=%s, date=%s",
             len(tickers), mode, currency, date)

    # Build workbook
    build_workbook(xlsx_path, tickers, metrics, date, curr_opt)

    # Launch Excel and process
    session = launch_excel_isolated(xlsx_path)
    try:
        excel = session.excel
        log.info("Connected to Excel %s", excel.Version)
        ws_com = session.workbook.Sheets(1)

        # Trigger refresh
        try:
            excel.Run("SNLXLAddin.xla!RefreshSheet")
            log.info("RefreshSheet executed")
        except Exception as e:
            log.warning("RefreshSheet warning: %s", e)

        # Poll for completion
        poll_for_completion(ws_com, len(tickers), len(metrics), max_wait)

        # Read results
        log.info("Reading data...")
        results, resolution_notes = read_results(ws_com, len(tickers), metrics)
    finally:
        close_session(session, save=False, delete_workbook=True)

    # Build table
    if mode == "lease-adjusted":
        df = build_table_lease_adjusted(results, lease_adjust_ntm)
    else:
        df = build_table_excluding_leases(results)

    # Convert to JSON-serializable dict
    data = dataframe_to_json(df, mode, currency, date, lease_adjust_ntm)
    data["resolution_notes"] = resolution_notes

    log.info("Comp pull complete: %d companies", len(data["companies"]))
    return data
