"""Formatted XLSX export for comp tables.

Renders the structured output of :func:`capiq_excel.engines.comps.run_comps`
into a presentation-ready workbook that matches the firm's "Comp Set" template
(Lato font; navy / teal / gold palette; title band, "As at" subtitle, wrapped
column headers, optional sector group sections, per-group + overall
Average/Median rows, and live derived formulas with ``NM`` guards).

The set of columns is configurable: callers pass an ordered list of column keys
(see :data:`AVAILABLE_KEYS` / :func:`list_columns`) and/or extra-metric specs.
Derived columns (margins, EV multiples, ND/EBITDA) are written as Excel
formulas referencing the raw input cells when those inputs are present, so the
deliverable recalculates if a number is edited; if a required input column was
dropped, the column falls back to a static value.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from datetime import datetime
from typing import Callable, Optional

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger("capiq_excel")

# ── Palette (resolved from the firm template theme) ────────────────────────
WHITE = "FFFFFFFF"
BLACK = "FF000000"
NAVY = "FF0A1748"   # column-header + overall-stat fill
TEAL = "FF3091BA"   # group-header fill
GOLD = "FFFAD043"   # per-group stat fill
GRAY = "FFBFBFBF"   # subtitle fill
GREEN = "FF00B050"  # ticker text

FONT_NAME = "Lato"

# ── Number formats ─────────────────────────────────────────────────────────
MONEY_FMT = r'_(* #,##0_);_(* \(#,##0\);_(* "-"??_);_(@_)'
PCT_FMT = r'0.0%;(0.0%)'
MULT_FMT = r'#,##0.0"x";(#,##0.0"x")'
COUNT_FMT = r'#,##0'


def _num(v) -> bool:
    """True if v is a usable finite number."""
    if v is None:
        return False
    if isinstance(v, str):
        return False
    try:
        return not np.isnan(float(v))
    except (TypeError, ValueError):
        return False


# ── Column model ───────────────────────────────────────────────────────────

@dataclass(frozen=True)
class XlsxColumn:
    key: str            # selection id (for the columns= argument)
    header: str         # display header text
    ctype: str          # "name" | "text" | "money" | "pct" | "mult" | "count"
    block: str          # vertical-separator group
    data_key: Optional[str] = None   # key into the company dict (value columns)
    summary: bool = False            # include in Average / Median rows
    derived: Optional[str] = None    # key into DERIVED (formula columns)
    width: float = 12.5


def _num_fmt(ctype: str) -> str:
    return {"money": MONEY_FMT, "pct": PCT_FMT, "mult": MULT_FMT,
            "count": COUNT_FMT}.get(ctype, "General")


# Derived columns: (required input data_keys, formula builder, value-fallback key).
# The formula builder receives a dict {data_key: "<col><row>"} of cell refs.
def _f_margin(c):
    return f'=IFERROR({c["LTM EBITDA"]}/{c["LTM Rev"]},"NM")'

def _f_ev(num_key):
    def build(c):
        n, d = c["TEV"], c[num_key]
        return f'=IFERROR(IF({n}/{d}>0,{n}/{d},"NM"),"NM")'
    return build

def _f_nd_ebitda(c):
    n, d = c["Net Debt"], c["LTM EBITDA"]
    return f'=IFERROR(IF({n}/{d}>0,{n}/{d},"NM"),"NM")'

DERIVED: dict[str, tuple[list[str], Callable[[dict], str], str]] = {
    "ebitda_margin": (["LTM EBITDA", "LTM Rev"], _f_margin, "EBITDA Margin %"),
    "ev_rev":        (["TEV", "LTM Rev"], _f_ev("LTM Rev"), "EV/Rev"),
    "ev_ltm_ebitda": (["TEV", "LTM EBITDA"], _f_ev("LTM EBITDA"), "EV/LTM EBITDA"),
    "ev_ntm_ebitda": (["TEV", "NTM EBITDA"], _f_ev("NTM EBITDA"), "EV/NTM EBITDA"),
    "nd_ebitda":     (["Net Debt", "LTM EBITDA"], _f_nd_ebitda, "ND/EBITDA"),
}


def _col(key, header, ctype, block, **kw) -> XlsxColumn:
    return XlsxColumn(key=key, header=header, ctype=ctype, block=block, **kw)


# Full registry of selectable columns per mode (includes optional columns not
# shown by default).  Order here is irrelevant; DEFAULT_KEYS sets the order.
def _registry(mode: str) -> dict[str, XlsxColumn]:
    lease = mode == "lease-adjusted"
    cols = [
        _col("company", "Company", "name", "cap", data_key="Company", width=44),
        _col("gaap", "GAAP", "text", "cap", data_key="GAAP", width=11),
        _col("mkt_cap", "Market Cap", "money", "cap", data_key="Mkt Cap", width=13),
        _col("total_debt", "Total Debt", "money", "cap", data_key="Total Debt", width=12.5),
        _col("leases", "Leases\n(US GAAP)", "money", "cap", data_key="Leases", width=12),
        _col("cash", "Cash and\nEquiv.", "money", "cap", data_key="Cash", width=12),
        _col("net_debt", "Net Debt", "money", "cap", data_key="Net Debt", width=12.5),
        _col("tev", "Enterprise\nValue", "money", "cap", data_key="TEV", width=13),
        _col("ltm_rev", "LTM\nRevenue", "money", "ops", data_key="LTM Rev", width=12.5),
        _col("ntm_rev", "NTM\nRevenue", "money", "ops", data_key="NTM Rev", width=12.5),
        _col("ltm_ebitda", "LTM\nEBITDA", "money", "ops", data_key="LTM EBITDA", width=12.5),
        _col("ntm_ebitda", "NTM\nEBITDA", "money", "ops", data_key="NTM EBITDA", width=12.5),
        _col("ebitda_margin", "EBITDA\nMargin %", "pct", "ops", derived="ebitda_margin", summary=True, width=11),
        _col("gross_margin", "Gross\nMargin %", "pct", "ops", data_key="Gross Margin %", summary=True, width=11),
        _col("capex", "CapEx", "money", "ops", data_key="CapEx", width=11),
        _col("ev_rev", "EV/\nRevenue", "mult", "evm", derived="ev_rev", summary=True, width=11),
        _col("ev_ltm_ebitda", "EV/\nLTM EBITDA", "mult", "evm", derived="ev_ltm_ebitda", summary=True, width=11.5),
        _col("ev_ntm_ebitda", "EV/\nNTM EBITDA", "mult", "evm", derived="ev_ntm_ebitda", summary=True, width=11.5),
        _col("pbv", "Price to\nBook", "mult", "othm", data_key="P/BV", summary=True, width=10.5),
        _col("nd_ebitda", "Net Debt/\nEBITDA", "mult", "othm", derived="nd_ebitda", summary=True, width=11),
    ]
    if lease:
        cols.insert(11, _col("lease_adj", "Lease\nAdj.", "money", "ops", data_key="Lease Adj", width=10))
    return {c.key: c for c in cols}


# Default visible columns (matches the firm template), per mode.
_DEFAULTS_LEASE = [
    "company", "mkt_cap", "total_debt", "leases", "cash", "net_debt", "tev",
    "ltm_rev", "ntm_rev", "ltm_ebitda", "lease_adj", "ntm_ebitda", "ebitda_margin",
    "ev_rev", "ev_ltm_ebitda", "ev_ntm_ebitda", "pbv", "nd_ebitda",
]
_DEFAULTS_EXCL = [
    "company", "mkt_cap", "total_debt", "leases", "cash", "net_debt", "tev",
    "ltm_rev", "ntm_rev", "ltm_ebitda", "ntm_ebitda", "ebitda_margin",
    "ev_rev", "ev_ltm_ebitda", "ev_ntm_ebitda", "pbv", "nd_ebitda",
]

# Block order for laying out vertical separators.
_BLOCK_ORDER = ["cap", "ops", "evm", "othm", "extra"]

AVAILABLE_KEYS = sorted(set(_registry("lease-adjusted")) | set(_registry("excluding-leases")))


def list_columns(mode: str = "lease-adjusted") -> list[tuple[str, str, bool]]:
    """Return (key, header, is_default) for every selectable column in a mode."""
    reg = _registry(mode)
    defaults = set(_DEFAULTS_LEASE if mode == "lease-adjusted" else _DEFAULTS_EXCL)
    rows = []
    for key, c in reg.items():
        rows.append((key, c.header.replace("\n", " "), key in defaults))
    return sorted(rows, key=lambda r: (not r[2], r[0]))


def _extra_column(spec) -> XlsxColumn:
    ctype = {"pct": "pct", "dollar": "money", "mult": "mult", "count": "count"}.get(spec.fmt, "money")
    return XlsxColumn(key=spec.key, header=spec.label, ctype=ctype, block="extra",
                      data_key=spec.label, summary=spec.is_summary,
                      width=12 if ctype == "money" else 11)


def build_columns(mode: str, selected_keys: Optional[list[str]], extra_specs) -> list[XlsxColumn]:
    """Resolve the ordered column list from selection + extras."""
    reg = _registry(mode)
    extra_cols = [_extra_column(s) for s in (extra_specs or [])]
    available = dict(reg)
    for ec in extra_cols:
        available[ec.key] = ec

    if selected_keys:
        # Keys valid in *some* mode (incl. extras). Keys outside this set are
        # typos and raise; keys known elsewhere but not in this mode (e.g.
        # lease_adj on the excluding-leases tab) are silently skipped so the
        # same --columns list can drive both tabs.
        known_anywhere = set(AVAILABLE_KEYS) | {ec.key for ec in extra_cols}
        unknown = [k for k in selected_keys if k not in known_anywhere]
        if unknown:
            raise ValueError(
                f"Unknown column key(s): {', '.join(unknown)}. "
                f"Available: {', '.join(sorted(known_anywhere))}"
            )
        cols = [available[k] for k in selected_keys if k in available]
        if not any(c.key == "company" for c in cols):
            cols.insert(0, available["company"])
    else:
        default = _DEFAULTS_LEASE if mode == "lease-adjusted" else _DEFAULTS_EXCL
        cols = [reg[k] for k in default] + extra_cols
    return cols


# ── Styling helpers ──────────────────────────────────────────────────────────

THIN = Side(style="thin", color=NAVY)


def _font(bold=False, italic=False, color=BLACK, size=11):
    return Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=color)


def _fill(rgb):
    return PatternFill(fill_type="solid", fgColor=rgb)


def _border(left=False, right=False, top=False, bottom=False):
    return Border(
        left=THIN if left else Side(),
        right=THIN if right else Side(),
        top=THIN if top else Side(),
        bottom=THIN if bottom else Side(),
    )


def _date_long(date_str: str) -> str:
    try:
        dt = datetime.strptime(date_str, "%m/%d/%Y")
        return f"{dt.strftime('%B')} {dt.day}, {dt.year}"
    except (ValueError, TypeError):
        return date_str


# ── Main writer ──────────────────────────────────────────────────────────────

def _sheet_title(mode: str) -> str:
    return "Lease Adjusted" if mode == "lease-adjusted" else "Excluding Leases"


def write_comps_xlsx(
    datasets,
    path: str,
    *,
    columns: Optional[list[str]] = None,
    groups: Optional[dict] = None,
    extra_specs=None,
    active_mode: Optional[str] = None,
) -> str:
    """Write a formatted comp-table workbook to *path*.

    Parameters
    ----------
    datasets : dict | list[dict]
        One :func:`capiq_excel.engines.comps.run_comps` output (single tab) or
        a list of them (one tab each, in order — e.g. a lease-adjusted dataset
        followed by an excluding-leases dataset to mirror the firm template).
    path : str
        Destination .xlsx path.
    columns : list[str], optional
        Ordered column keys (see :func:`list_columns`). Defaults to the
        template's standard column set for each tab's mode. Keys valid in one
        mode but not another (e.g. ``lease_adj``) are silently skipped on tabs
        where they don't apply.
    groups : dict, optional
        ``{label: [tickers]}`` sector grouping (US exchange prefixes are
        stripped to match stored tickers).
    extra_specs : list[ExtraMetricSpec], optional
        Extra-metric specs (so extra columns get correct headers / formats).
    active_mode : str, optional
        Which tab to make active on open ("lease-adjusted" / "excluding-leases").
    """
    if isinstance(datasets, dict):
        datasets = [datasets]

    wb = Workbook()
    for i, data in enumerate(datasets):
        ws = wb.active if i == 0 else wb.create_sheet()
        _write_sheet(ws, data, columns=columns, groups=groups, extra_specs=extra_specs)

    if active_mode in ("lease-adjusted", "excluding-leases"):
        target = _sheet_title(active_mode)
        if target in wb.sheetnames:
            wb.active = wb.sheetnames.index(target)

    wb.save(path)
    log.info("Wrote formatted comp workbook: %s (%d tab%s)",
             path, len(datasets), "" if len(datasets) == 1 else "s")
    return path


def _write_sheet(ws, data: dict, *, columns=None, groups=None, extra_specs=None) -> None:
    """Render one comp-table tab into an existing worksheet."""
    mode = data.get("mode", "lease-adjusted")
    currency = (data.get("currency") or "USD").upper()
    date = data.get("date", "")
    lease_adjust_ntm = data.get("lease_adjust_ntm", True)
    companies = list(data.get("companies", []))

    cols = build_columns(mode, columns, extra_specs)

    ws.title = _sheet_title(mode)
    ws.sheet_view.showGridLines = False

    GUTTER, TICKER = 1, 2
    c0 = 3                       # first table column (Company)
    cN = c0 + len(cols) - 1      # last table column
    col_letter = {c.key: get_column_letter(c0 + i) for i, c in enumerate(cols)}
    # data_key -> column letter, for derived-formula references
    ref_map = {c.data_key: col_letter[c.key] for c in cols if c.data_key}

    # First column of each present block gets a left separator.
    present_blocks = [b for b in _BLOCK_ORDER if any(c.block == b for c in cols)]
    block_first_idx = {}
    for i, c in enumerate(cols):
        if c.block not in block_first_idx:
            block_first_idx[c.block] = i

    def left_sep(i):  # vertical separator before column index i
        return cols[i].block in block_first_idx and block_first_idx[cols[i].block] == i

    # ── Column widths ──
    ws.column_dimensions[get_column_letter(GUTTER)].width = 2.6
    ws.column_dimensions[get_column_letter(TICKER)].width = 13
    for i, c in enumerate(cols):
        ws.column_dimensions[get_column_letter(c0 + i)].width = c.width

    r = 2  # leave row 1 as a top margin

    # ── Title band ──
    title = "Comparable Companies Analysis"
    ws.cell(r, c0, title)
    for ci in range(c0, cN + 1):
        cell = ws.cell(r, ci)
        cell.font = _font(bold=True, color=WHITE)
        cell.fill = _fill(BLACK)
        cell.border = _border(left=(ci == c0), right=(ci == cN), top=True, bottom=True)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=cN)
    ws.row_dimensions[r].height = 20
    r += 1

    # ── Subtitle ("As at ...") ──
    subtitle = f"As at {_date_long(date)}        Figures in {currency}$ millions"
    ws.cell(r, c0, subtitle)
    for ci in range(c0, cN + 1):
        cell = ws.cell(r, ci)
        cell.font = _font(bold=True, italic=True, color=BLACK)
        cell.fill = _fill(GRAY)
        cell.border = _border(left=(ci == c0), right=(ci == cN), bottom=True)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=cN)
    ws.row_dimensions[r].height = 16
    r += 1

    # ── Column header row ──
    header_row = r
    for i, c in enumerate(cols):
        ci = c0 + i
        cell = ws.cell(header_row, ci, c.header)
        cell.font = _font(bold=True, color=WHITE)
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(
            horizontal="left" if c.ctype == "name" else "right",
            vertical="center", wrap_text=True,
        )
        cell.border = _border(left=left_sep(i) or ci == c0, right=(ci == cN),
                              top=True, bottom=True)
    ws.row_dimensions[header_row].height = 31.5
    r += 1

    def write_company_row(company, row):
        # Ticker (floating, green, no border)
        tk = ws.cell(row, TICKER, company.get("Ticker", ""))
        tk.font = _font(color=GREEN)
        tk.alignment = Alignment(horizontal="right")
        for i, c in enumerate(cols):
            ci = c0 + i
            cell = ws.cell(row, ci)
            value, is_formula = _cell_content(c, company, ref_map, row)
            if value is not None:
                cell.value = value
            cell.font = _font()
            cell.number_format = _num_fmt(c.ctype)
            cell.alignment = Alignment(
                horizontal="left" if c.ctype in ("name", "text") else "right",
                vertical="center",
            )
            cell.border = _border(left=left_sep(i) or ci == c0, right=(ci == cN))

    def write_stat_row(label, func, spans, row, fill, txt_color, top=False, bottom=False):
        """spans: list of (first_row, last_row) of *company* rows to aggregate."""
        for i, c in enumerate(cols):
            ci = c0 + i
            cell = ws.cell(row, ci)
            cell.fill = _fill(fill)
            cell.font = _font(bold=True, color=txt_color)
            cell.alignment = Alignment(
                horizontal="left" if c.ctype in ("name", "text") else "right",
                vertical="center",
            )
            cell.border = _border(left=left_sep(i) or ci == c0, right=(ci == cN),
                                  top=top, bottom=bottom)
            if i == 0:
                cell.value = label
            elif c.summary:
                letter = col_letter[c.key]
                rng = ",".join(f"{letter}{a}:{letter}{b}" for a, b in spans)
                cell.value = f'=IFERROR({func}({rng}),"NM")'
                cell.number_format = _num_fmt(c.ctype)

    # ── Body: grouped or flat ──
    def resolve_group(tickers):
        from capiq_excel.engines.comps import strip_us_exchange_prefix
        stripped = {strip_us_exchange_prefix(t) for t in tickers}
        return [co for co in companies if co.get("Ticker") in stripped]

    def write_group(label, members, fill_stat, txt_stat):
        nonlocal r
        # teal group header
        for ci in range(c0, cN + 1):
            cell = ws.cell(r, ci, label if ci == c0 else None)
            cell.font = _font(bold=True, color=WHITE)
            cell.fill = _fill(TEAL)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = _border(left=(ci == c0), right=(ci == cN), top=True, bottom=True)
        r += 1
        first = r
        for co in members:
            write_company_row(co, r)
            r += 1
        span = (first, r - 1)
        write_stat_row(f"Average {label}", "AVERAGE", [span], r, fill_stat, txt_stat, top=True)
        r += 1
        write_stat_row(f"Median {label}", "MEDIAN", [span], r, fill_stat, txt_stat, bottom=True)
        r += 1
        return span

    if groups:
        used = set()
        spans = []
        for label, tickers in groups.items():
            members = resolve_group(tickers)
            if not members:
                continue
            for co in members:
                used.add(id(co))
            spans.append(write_group(label, members, GOLD, BLACK))
        leftovers = [co for co in companies if id(co) not in used]
        if leftovers:
            spans.append(write_group("Other", leftovers, GOLD, BLACK))
        # overall — aggregate only company rows (skip per-group stat rows)
        write_stat_row("Overall Average", "AVERAGE", spans, r, NAVY, WHITE, top=True)
        r += 1
        write_stat_row("Overall Median", "MEDIAN", spans, r, NAVY, WHITE, bottom=True)
        r += 1
    else:
        first = r
        for co in companies:
            write_company_row(co, r)
            r += 1
        span = [(first, r - 1)]
        write_stat_row("Average", "AVERAGE", span, r, NAVY, WHITE, top=True)
        r += 1
        write_stat_row("Median", "MEDIAN", span, r, NAVY, WHITE, bottom=True)
        r += 1

    # ── Footnotes ──
    r += 1
    notes = []
    if mode == "lease-adjusted" and lease_adjust_ntm:
        notes.append("* NTM EBITDA lease-adjusted: TTM lease expense added to consensus "
                     "NTM EBITDA for US GAAP reporters as a proxy.")
    elif mode == "lease-adjusted":
        notes.append("NTM EBITDA shown as straight consensus (lease adjustment disabled).")
    else:
        notes.append("TEV and debt exclude operating leases; EBITDA excludes lease adjustment.")
    notes.append("Source: S&P Capital IQ. \"NM\" = not meaningful.")
    for note in notes:
        cell = ws.cell(r, c0, note)
        cell.font = _font(italic=True, color="FF808080", size=8)
        r += 1

    # Freeze header + identity columns
    ws.freeze_panes = ws.cell(header_row + 1, c0)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    log.info("Rendered '%s' tab: %d companies, %d columns", ws.title, len(companies), len(cols))


def _cell_content(col: XlsxColumn, company: dict, ref_map: dict, row: int):
    """Return (value_or_formula, is_formula) for a data cell."""
    if col.key == "company":
        from capiq_excel.engines.comps import is_error_value
        name = company.get("Company")
        if not name or is_error_value(str(name)):
            name = company.get("Ticker", "")
        name = str(name)
        if company.get("Lease Adj Applied"):
            name += " *"
        return name, False

    if col.derived:
        needs, build, fallback_key = DERIVED[col.derived]
        if all(n in ref_map for n in needs):
            refs = {n: f"{ref_map[n]}{row}" for n in needs}
            return build(refs), True
        # fallback to a static value
        if col.ctype == "pct":
            v = company.get(fallback_key)
            return (float(v) / 100.0 if _num(v) else None), False
        v = company.get(fallback_key)
        if not _num(v):
            return None, False
        return (float(v) if float(v) > 0 else "NM"), False

    # plain value column
    v = company.get(col.data_key)
    if col.ctype == "pct":
        return (float(v) / 100.0 if _num(v) else None), False
    if col.ctype in ("money", "mult", "count"):
        return (float(v) if _num(v) else None), False
    # text / name
    return (str(v) if v not in (None, "") and not (isinstance(v, float) and np.isnan(v)) else None), False
