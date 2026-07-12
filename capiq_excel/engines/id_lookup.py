"""Lightweight identifier resolution via CIQRANGEA + CIQ company name lookup.

Builds a small workbook with just ID resolution formulas, launches Excel,
reads results. Much faster than a full comp pull (~10s vs ~40s).
"""
from __future__ import annotations

import logging
import os
import tempfile
import time

import pythoncom
from openpyxl import Workbook

from capiq_excel.engines.comps import is_error_value

log = logging.getLogger("capiq_excel")


def build_id_workbook(path: str, identifiers: list[str]) -> None:
    """Build a workbook with CIQRANGEA ID lookup + CIQ company name per identifier.

    Layout:
      Col A: identifier (raw input)
      Col B: =CIQRANGEA(id, "IQ_COMPANY_ID_QUICK_MATCH", 1, 1)  (formula)
      Col C: (blank — CIQRANGEA spill result = resolved IQ ID)
      Col D: =CIQ(spill_ref, "IQ_COMPANY_NAME")  (name from resolved ID)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Lookup"

    ws.cell(row=1, column=1, value="Input")
    ws.cell(row=1, column=2, value="CIQRANGEA_Formula")
    ws.cell(row=1, column=3, value="Resolved_IQ_ID")
    ws.cell(row=1, column=4, value="Company_Name")

    for idx, ident in enumerate(identifiers):
        row = idx + 2
        # CIQRANGEA doesn't handle EXCHANGE:TICKER format — strip prefix
        lookup_id = ident.split(":", 1)[1] if ":" in ident else ident
        ws.cell(row=row, column=1, value=ident)
        ws.cell(row=row, column=2,
                value=f'=CIQRANGEA("{lookup_id}","IQ_COMPANY_ID_QUICK_MATCH",1,1)')
        # Col C is blank — CIQRANGEA result spills here
        # Company name from resolved IQ ID (references spill cell)
        ws.cell(row=row, column=4,
                value=f'=CIQ($C${row},"IQ_COMPANY_NAME")')

    wb.save(path)
    log.info("Created ID lookup workbook: %s (%d identifiers)", path, len(identifiers))


def run_id_lookup(identifiers: list[str], max_wait: int = 60) -> dict:
    """Resolve identifiers to CIQ IQ IDs and company names."""
    if not identifiers:
        raise ValueError("identifiers must not be empty")
    if max_wait <= 0:
        raise ValueError("max_wait must be greater than zero")

    from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

    pythoncom.CoInitialize()
    try:
        with tempfile.TemporaryDirectory(prefix="capiq-lookup-") as temp_dir:
            xlsx_path = os.path.join(temp_dir, "lookup.xlsx")
            build_id_workbook(xlsx_path, identifiers)

            session = launch_excel_isolated(xlsx_path)
            try:
                excel = session.excel
                log.info("Connected to Excel %s for ID lookup", excel.Version)
                ws_com = session.workbook.Sheets(1)

                log.info("Calling RefreshSheet for ID lookup...")
                try:
                    pythoncom.PumpWaitingMessages()
                    excel.Run("SNLXLAddin.xla!RefreshSheet")
                    log.info("RefreshSheet executed for ID lookup")
                except Exception as e:
                    log.warning("RefreshSheet warning: %s", e)
                pythoncom.PumpWaitingMessages()

                n = len(identifiers)
                start = time.monotonic()
                while True:
                    elapsed = time.monotonic() - start
                    if elapsed > max_wait:
                        log.warning("ID lookup timeout after %ds", max_wait)
                        break

                    if n == 1:
                        spill_vals = ((ws_com.Cells(2, 3).Value,),)
                    else:
                        spill_vals = ws_com.Range(
                            ws_com.Cells(2, 3), ws_com.Cells(1 + n, 3)
                        ).Value

                    pending = 0
                    for row_tuple in spill_vals:
                        value = row_tuple[0]
                        if value is None:
                            pending += 1
                        elif isinstance(value, str) and value.strip().upper() in (
                            "#PEND", "#REFRESH"
                        ):
                            pending += 1
                        elif isinstance(value, int) and value == -2146826259:
                            pending += 1

                    log.debug(
                        "ID lookup %ds: %d/%d settled, %d pending",
                        elapsed, n - pending, n, pending,
                    )
                    if pending == 0:
                        break

                    pythoncom.PumpWaitingMessages()
                    time.sleep(2)

                if n == 1:
                    result_vals = ((
                        ws_com.Cells(2, 3).Value,
                        ws_com.Cells(2, 4).Value,
                    ),)
                else:
                    result_vals = ws_com.Range(
                        ws_com.Cells(2, 3), ws_com.Cells(1 + n, 4)
                    ).Value

                results = []
                for idx, ident in enumerate(identifiers):
                    iq_id = result_vals[idx][0]
                    company_name = result_vals[idx][1]
                    if is_error_value(iq_id):
                        results.append({
                            "input": ident,
                            "iq_id": None,
                            "company_name": None,
                            "status": "FAILED",
                        })
                    else:
                        name = (
                            company_name
                            if isinstance(company_name, str)
                            and not is_error_value(company_name)
                            else None
                        )
                        results.append({
                            "input": ident,
                            "iq_id": str(iq_id),
                            "company_name": name,
                            "status": "OK",
                        })
            finally:
                close_session(session, save=False, delete_workbook=False)

        log.info(
            "ID lookup complete: %d/%d resolved",
            sum(1 for result in results if result["status"] == "OK"),
            len(results),
        )
        return {"results": results}
    finally:
        pythoncom.CoUninitialize()
