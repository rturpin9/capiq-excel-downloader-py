"""
Investment Banking Comparable Companies Table via Capital IQ.

Downloads fundamental data for a set of companies using CIQ() formulas,
calculates valuation multiples, and outputs a formatted comp table.
"""
import time
import os
import pandas as pd
import numpy as np
from openpyxl import Workbook

import pythoncom
from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

pythoncom.CoInitialize()

# -- Companies -----------------------------------------------------------
# Each entry: (display_name, primary_ticker, fallback_ticker_or_None)
COMPANIES = [
    ("Descartes Systems",       "DSGX",     None),
    ("IDEX Corp",               "IEX",      None),
    ("Computer Modelling Grp",  "CMDXF",    "CMG.TO"),
    ("Manhattan Associates",    "MANH",     None),
    ("WiseTech Global",         "WTC",      None),
    ("Roper Technologies",      "ROP",      None),
    ("Constellation Software",  "CNSWF",    "CSU.TO"),
]

# -- Metrics -------------------------------------------------------------
# Each entry: (display_label, primary_mnemonic, [fallback_mnemonics...])
METRICS = [
    ("Company Name",    "IQ_COMPANY_NAME",      []),
    ("Close Price",     "IQ_CLOSEPRICE",        []),
    ("Market Cap",      "IQ_MARKETCAP",         []),
    ("Ent. Value",      "IQ_TEV",               ["IQ_ENTERPRISE_VALUE", "IQ_EV"]),
    ("LTM Revenue",     "IQ_TOTAL_REV",         []),
    ("LTM EBITDA",      "IQ_EBITDA",            []),
    ("NTM Revenue",     "IQ_MEAN_REVENUE_EST",  ["IQ_REVENUE_EST", "IQ_EST_REV_MEAN", "IQ_CONSENSUS_REVENUE"]),
    ("NTM EBITDA",      "IQ_MEAN_EBITDA_EST",   ["IQ_EBITDA_EST", "IQ_EST_EBITDA_MEAN", "IQ_CONSENSUS_EBITDA"]),
]


def build_workbook(path: str):
    """Create an XLSX with CIQ() formulas in a grid layout."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comp"

    # -- Primary grid ----------------------------------------------------
    # Row 1: headers
    ws.cell(row=1, column=1, value="Ticker")
    for col_idx, (label, _, _) in enumerate(METRICS, start=2):
        ws.cell(row=1, column=col_idx, value=label)

    # Rows 2+: one per company (primary ticker)
    for row_idx, (name, ticker, _) in enumerate(COMPANIES, start=2):
        ws.cell(row=row_idx, column=1, value=ticker)
        for col_idx, (_, mnemonic, _) in enumerate(METRICS, start=2):
            ws.cell(row=row_idx, column=col_idx,
                    value=f'=CIQ("{ticker}","{mnemonic}")')

    primary_rows = len(COMPANIES)

    # -- Fallback section: alternate tickers -----------------------------
    # For companies with a fallback ticker, add another row
    fb_start = primary_rows + 3  # leave a gap
    ws.cell(row=fb_start, column=1, value="-- Fallback Tickers --")
    fb_row = fb_start + 1
    fallback_map = {}  # (company_idx, metric_idx) -> (sheet_row, sheet_col)

    for comp_idx, (name, _, fallback) in enumerate(COMPANIES):
        if fallback is None:
            continue
        ws.cell(row=fb_row, column=1, value=fallback)
        for col_idx, (_, mnemonic, _) in enumerate(METRICS, start=2):
            ws.cell(row=fb_row, column=col_idx,
                    value=f'=CIQ("{fallback}","{mnemonic}")')
            fallback_map[(comp_idx, col_idx - 2)] = (fb_row, col_idx)
        fb_row += 1

    # -- Fallback section: alternate metric names ------------------------
    fb_row += 1
    ws.cell(row=fb_row, column=1, value="-- Fallback Metrics --")
    fb_row += 1
    metric_fb_cells = {}  # (comp_idx, metric_label) -> [(row, col), ...]

    for met_idx, (label, _, fallbacks) in enumerate(METRICS):
        if not fallbacks:
            continue
        for fb_mnemonic in fallbacks:
            for comp_idx, (_, ticker, _) in enumerate(COMPANIES):
                ws.cell(row=fb_row, column=1, value=f"{ticker}|{fb_mnemonic}")
                ws.cell(row=fb_row, column=2,
                        value=f'=CIQ("{ticker}","{fb_mnemonic}")')
                key = (comp_idx, label)
                metric_fb_cells.setdefault(key, []).append((fb_row, 2))
                fb_row += 1

    wb.save(path)
    return fallback_map, metric_fb_cells


def read_grid(ws_com, num_companies, num_metrics):
    """Read the primary grid values into a 2D dict."""
    data = {}
    for comp_idx in range(num_companies):
        row = comp_idx + 2
        ticker = ws_com.Cells(row, 1).Value
        row_data = {}
        for met_idx in range(num_metrics):
            col = met_idx + 2
            val = ws_com.Cells(row, col).Value
            label = METRICS[met_idx][0]
            row_data[label] = val
        data[comp_idx] = row_data
    return data


def resolve_fallbacks(ws_com, data, fallback_map, metric_fb_cells):
    """Fill in values from fallback tickers/metrics where primary failed."""
    error_strings = {
        "#INVALID COMPANY ID", "#INVALID METRIC NAME",
        "#INVALID FUNCTION PARAMETER", "#ERROR", "#NAME",
        "#OUTSIDE SUBSCRIPTION", "CIQRANGEA", "CIQRANGE",
        "(INVALID IDENTIFIER)", "(INVALID FORMULA NAME)",
        "(ERROR)", "INVALID", "#PEND",
    }

    def is_error(val):
        if val is None:
            return True
        if isinstance(val, str) and any(e in val.upper() for e in error_strings):
            return True
        if isinstance(val, (int, float)) and val < -2000000000:
            return True
        return False

    # Fallback tickers (international companies)
    for (comp_idx, met_idx), (fb_row, fb_col) in fallback_map.items():
        label = METRICS[met_idx][0]
        if is_error(data[comp_idx].get(label)):
            fb_val = ws_com.Cells(fb_row, fb_col).Value
            if not is_error(fb_val):
                data[comp_idx][label] = fb_val

    # Fallback metric names
    for (comp_idx, label), cells in metric_fb_cells.items():
        if is_error(data[comp_idx].get(label)):
            for fb_row, fb_col in cells:
                fb_val = ws_com.Cells(fb_row, fb_col).Value
                if not is_error(fb_val):
                    data[comp_idx][label] = fb_val
                    break

    return data


def safe_float(val):
    """Convert a cell value to float, returning NaN for errors."""
    if val is None:
        return np.nan
    if isinstance(val, (int, float)):
        if val < -2000000000:  # COM error code
            return np.nan
        return float(val)
    if isinstance(val, str):
        try:
            return float(val.replace(",", ""))
        except ValueError:
            return np.nan
    return np.nan


def build_comp_table(data):
    """Build a DataFrame comp table with derived multiples."""
    rows = []
    for comp_idx, (name, ticker, fallback) in enumerate(COMPANIES):
        d = data[comp_idx]
        company_name = d.get("Company Name", name)
        if isinstance(company_name, str) and len(company_name) > 2:
            display = company_name
        else:
            display = name

        price = safe_float(d.get("Close Price"))
        mkt_cap = safe_float(d.get("Market Cap"))
        ev = safe_float(d.get("Ent. Value"))
        ltm_rev = safe_float(d.get("LTM Revenue"))
        ltm_ebitda = safe_float(d.get("LTM EBITDA"))
        ntm_rev = safe_float(d.get("NTM Revenue"))
        ntm_ebitda = safe_float(d.get("NTM EBITDA"))

        # Derived multiples
        ev_rev_ltm = ev / ltm_rev if ltm_rev and not np.isnan(ev) and not np.isnan(ltm_rev) and ltm_rev != 0 else np.nan
        ev_rev_ntm = ev / ntm_rev if ntm_rev and not np.isnan(ev) and not np.isnan(ntm_rev) and ntm_rev != 0 else np.nan
        ev_ebitda_ltm = ev / ltm_ebitda if ltm_ebitda and not np.isnan(ev) and not np.isnan(ltm_ebitda) and ltm_ebitda != 0 else np.nan
        ev_ebitda_ntm = ev / ntm_ebitda if ntm_ebitda and not np.isnan(ev) and not np.isnan(ntm_ebitda) and ntm_ebitda != 0 else np.nan
        ebitda_margin = (ltm_ebitda / ltm_rev * 100) if ltm_rev and ltm_ebitda and not np.isnan(ltm_rev) and not np.isnan(ltm_ebitda) and ltm_rev != 0 else np.nan

        rows.append({
            "Company": display,
            "Ticker": ticker,
            "Price": price,
            "Mkt Cap ($M)": mkt_cap,
            "EV ($M)": ev,
            "LTM Rev ($M)": ltm_rev,
            "LTM EBITDA ($M)": ltm_ebitda,
            "NTM Rev ($M)": ntm_rev,
            "NTM EBITDA ($M)": ntm_ebitda,
            "EV/Rev LTM": ev_rev_ltm,
            "EV/Rev NTM": ev_rev_ntm,
            "EV/EBITDA LTM": ev_ebitda_ltm,
            "EV/EBITDA NTM": ev_ebitda_ntm,
            "EBITDA Margin %": ebitda_margin,
        })

    df = pd.DataFrame(rows)
    return df


def format_table(df):
    """Print a nicely formatted comp table to the console."""
    print("\n" + "=" * 120)
    print("  COMPARABLE COMPANIES ANALYSIS")
    print("=" * 120)

    # Column formatting
    fmt = {
        "Price": "${:,.2f}",
        "Mkt Cap ($M)": "${:,.0f}",
        "EV ($M)": "${:,.0f}",
        "LTM Rev ($M)": "${:,.1f}",
        "LTM EBITDA ($M)": "${:,.1f}",
        "NTM Rev ($M)": "${:,.1f}",
        "NTM EBITDA ($M)": "${:,.1f}",
        "EV/Rev LTM": "{:.1f}x",
        "EV/Rev NTM": "{:.1f}x",
        "EV/EBITDA LTM": "{:.1f}x",
        "EV/EBITDA NTM": "{:.1f}x",
        "EBITDA Margin %": "{:.1f}%",
    }

    # Header
    header_cols = ["Company"] + [c for c in df.columns if c not in ("Company", "Ticker")]
    header_widths = {"Company": 35}
    for c in header_cols:
        if c == "Company":
            continue
        header_widths[c] = max(len(c), 14)

    header_line = f"{'Company':<35}"
    for c in header_cols[1:]:
        header_line += f"  {c:>{header_widths[c]}}"
    print(f"\n{header_line}")
    print("-" * len(header_line))

    # Data rows
    for _, row in df.iterrows():
        # Truncate company name
        cname = str(row["Company"])[:33]
        ticker = row["Ticker"]
        line = f"{cname:<35}"
        for c in header_cols[1:]:
            val = row[c]
            w = header_widths[c]
            if pd.isna(val):
                line += f"  {'N/A':>{w}}"
            elif c in fmt:
                line += f"  {fmt[c].format(val):>{w}}"
            else:
                line += f"  {str(val):>{w}}"
        print(line)

    # Summary stats
    print("-" * len(header_line))
    numeric_cols = [c for c in header_cols[1:] if c in fmt]
    for stat_name, stat_fn in [("Mean", "mean"), ("Median", "median")]:
        line = f"{stat_name:<35}"
        for c in header_cols[1:]:
            w = header_widths[c]
            if c in fmt:
                vals = df[c].dropna()
                if len(vals) > 0:
                    s = getattr(vals, stat_fn)()
                    line += f"  {fmt[c].format(s):>{w}}"
                else:
                    line += f"  {'N/A':>{w}}"
            else:
                line += f"  {'':>{w}}"
        print(line)

    print("=" * len(header_line))


# -- Main ----------------------------------------------------------------
def main():
    xlsx_path = os.path.abspath("comp_data_temp.xlsx")
    csv_path = os.path.abspath("comp_table.csv")

    print("Building comp table workbook...")
    fallback_map, metric_fb_cells = build_workbook(xlsx_path)
    print(f"  Created: {xlsx_path}")
    print(f"  {len(COMPANIES)} companies x {len(METRICS)} metrics")

    # Launch isolated Excel instance (does NOT kill existing Excel windows)
    print("\nLaunching Excel (isolated instance)...")
    session = launch_excel_isolated(xlsx_path)
    excel = session.excel
    print(f"Connected to Excel {excel.Version}")
    ws_com = session.workbook.Sheets(1)

    # Trigger refresh
    print("Triggering RefreshSheet...")
    try:
        excel.Run("SNLXLAddin.xla!RefreshSheet")
        print("  RefreshSheet executed")
    except Exception as e:
        print(f"  RefreshSheet warning: {e}")

    # Poll for completion
    MAX_WAIT = 120
    POLL_INTERVAL = 5
    print(f"Waiting for formulas to evaluate (max {MAX_WAIT}s)...")
    start = time.monotonic()

    while True:
        elapsed = time.monotonic() - start
        if elapsed > MAX_WAIT:
            print(f"  Timeout after {MAX_WAIT}s -- proceeding with available data")
            break

        # Check a sample of cells for pending tokens
        pending = False
        resolved = 0
        for comp_idx in range(len(COMPANIES)):
            row = comp_idx + 2
            for met_idx in range(len(METRICS)):
                col = met_idx + 2
                val = ws_com.Cells(row, col).Value
                if isinstance(val, str) and val.upper() in ('#PEND', '#REFRESH'):
                    pending = True
                elif val is not None and not (isinstance(val, int) and val < -2000000000):
                    resolved += 1

        total = len(COMPANIES) * len(METRICS)
        if not pending and resolved > total * 0.5:
            print(f"  Data ready! {resolved}/{total} cells resolved after {elapsed:.0f}s")
            break

        if int(elapsed) % 15 == 0:
            print(f"  {elapsed:.0f}s: {resolved}/{total} resolved, pending={pending}")

        time.sleep(POLL_INTERVAL)

    # Read data
    print("\nReading data...")
    data = read_grid(ws_com, len(COMPANIES), len(METRICS))
    data = resolve_fallbacks(ws_com, data, fallback_map, metric_fb_cells)

    # Debug: show raw values
    print("\n-- Raw Values --")
    for comp_idx, (name, ticker, _) in enumerate(COMPANIES):
        d = data[comp_idx]
        vals = ", ".join(f"{k}={v!r}" for k, v in d.items() if k != "Company Name")
        print(f"  {ticker}: {vals}")

    # Close only our Excel instance (other Excel windows are untouched)
    close_session(session, save=False, delete_workbook=True)

    # Build and display table
    df = build_comp_table(data)
    format_table(df)

    # Save CSV
    df.to_csv(csv_path, index=False)
    print(f"\nCSV saved to: {csv_path}")


if __name__ == "__main__":
    main()
