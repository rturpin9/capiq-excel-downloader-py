"""Download daily DSGX close prices via Capital IQ and plot a 12-month chart."""
import time
import os
import re
import subprocess
import win32com.client
import pythoncom
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook

pythoncom.CoInitialize()

# ── Config ──────────────────────────────────────────────────────────────
TICKER = "DSGX"
CIQ_ID = "IQ35352"  # Resolved CIQ ID for Descartes Systems Group
BEGIN_DATE = "02/25/2025"
END_DATE = "02/25/2026"
CHART_PATH = os.path.abspath("dsgx_stock_chart.png")

# Regex to extract date from expanded CIQ formula: =CIQ("id", "metric", "date")
CIQ_FORMULA_DATE_RE = re.compile(r'=CIQ\("[^"]+",\s*"[^"]+",\s*"([^"]+)"\)')

# ── Step 1: Create workbook with CIQRANGE formula ───────────────────────
print(f"Creating workbook for {TICKER} ({CIQ_ID}) daily close prices...")
print(f"Date range: {BEGIN_DATE} to {END_DATE}")

wb = Workbook()
ws = wb.active
ws.title = "Sheet"

# Column A: Date column (financial data style) — label only
ws['A1'] = "Date"

# Column B: Close price via CIQRANGE market data formula (legacy syntax)
# This will expand vertically, with each cell becoming =CIQ(id, metric, date)
formula = f'=CIQRANGE("{CIQ_ID}", "IQ_CLOSEPRICE", "{BEGIN_DATE}", "{END_DATE}", , , , , "IQ_CLOSEPRICE")'
ws['B1'] = formula
print(f"Formula: {formula}")

test_path = os.path.abspath("dsgx_daily_prices.xlsx")
wb.save(test_path)

# ── Step 2: Launch Excel via subprocess (required for add-in initialization) ──
print("\nKilling any existing Excel instances...")
os.system('taskkill /f /im excel.exe 2>NUL')
time.sleep(5)

excel_exe = r"C:\Program Files\Microsoft Office\root\Office16\EXCEL.EXE"
print(f"Launching Excel with workbook...")
proc = subprocess.Popen([excel_exe, test_path])
print("Waiting 30s for add-ins to initialize...")
time.sleep(30)

# Connect via Running Object Table
excel = win32com.client.GetActiveObject("Excel.Application")
print(f"Connected to Excel {excel.Version}")

# ── Step 3: Trigger refresh and wait ─────────────────────────────────────
print("\nTriggering RefreshSheet...")
try:
    excel.Run("SNLXLAddin.xla!RefreshSheet")
    print("RefreshSheet executed successfully")
except Exception as e:
    print(f"RefreshSheet warning: {e}")

# Poll for completion
MAX_WAIT = 180  # 3 minutes
POLL_INTERVAL = 5
ws_com = excel.ActiveWorkbook.Sheets(1)

print(f"Waiting for formulas to evaluate (max {MAX_WAIT}s)...")
start = time.monotonic()
last_count = 0

while True:
    elapsed = time.monotonic() - start
    if elapsed > MAX_WAIT:
        print(f"  Timeout after {MAX_WAIT}s")
        break

    # Count populated data cells in column B (starting from row 2)
    data_count = 0
    pending = False
    for row in range(2, 400):
        val = ws_com.Cells(row, 2).Value
        if val is None:
            break
        data_count += 1
        if isinstance(val, str) and val.upper() in ('#PEND', '#REFRESH', 'CIQRANGE'):
            pending = True

    if data_count != last_count:
        print(f"  {elapsed:.0f}s: {data_count} data rows, pending={pending}")
        last_count = data_count

    if data_count > 10 and not pending:
        # Check a sample cell to confirm it has a real value
        sample = ws_com.Cells(10, 2).Value
        if isinstance(sample, (int, float)) and sample > 0:
            print(f"  Data ready! {data_count} rows after {elapsed:.0f}s")
            break

    time.sleep(POLL_INTERVAL)

# ── Step 4: Extract data (dates from formulas, values from cells) ────────
print("\nExtracting data...")
dates = []
prices = []

for row in range(2, 500):
    val = ws_com.Cells(row, 2).Value
    if val is None:
        break

    # Get the formula — in compat mode, CIQRANGE expands to individual CIQ() calls
    try:
        formula = ws_com.Cells(row, 2).Formula
    except:
        formula = ""

    # Extract date from formula
    date_match = CIQ_FORMULA_DATE_RE.match(formula)
    if date_match:
        date_str = date_match.group(1)
    else:
        date_str = None

    if isinstance(val, (int, float)) and date_str:
        dates.append(date_str)
        prices.append(float(val))
    elif isinstance(val, str) and val.startswith('#'):
        print(f"  Row {row}: error = {val}")

print(f"Extracted {len(prices)} data points")

# ── Step 5: Close Excel and clean up ─────────────────────────────────────
excel.ActiveWorkbook.Close(SaveChanges=False)
excel.Quit()
os.remove(test_path)

# ── Step 6: Build DataFrame and plot ─────────────────────────────────────
if len(prices) < 5:
    print("ERROR: Not enough data points for a chart")
    exit(1)

df = pd.DataFrame({"Date": pd.to_datetime(dates), "Close": prices})
df = df.sort_values("Date")
df = df.set_index("Date")

print(f"\nData summary:")
print(f"  Period: {df.index[0].date()} to {df.index[-1].date()}")
print(f"  Points: {len(df)}")
print(f"  Min:    ${df['Close'].min():.2f}")
print(f"  Max:    ${df['Close'].max():.2f}")
print(f"  Latest: ${df['Close'].iloc[-1]:.2f}")

# Plot
fig, ax = plt.subplots(figsize=(14, 6))
ax.plot(df.index, df["Close"], color="#1f77b4", linewidth=1.2)
ax.fill_between(df.index, df["Close"], alpha=0.1, color="#1f77b4")
ax.set_title(f"Descartes Systems Group (DSGX) — Daily Close Price", fontsize=14, fontweight="bold")
ax.set_ylabel("Price (USD)", fontsize=12)
ax.set_xlabel("")
ax.grid(True, alpha=0.3)
ax.tick_params(axis='x', rotation=30)
fig.tight_layout()
fig.savefig(CHART_PATH, dpi=150, bbox_inches="tight")
print(f"\nChart saved to: {CHART_PATH}")
