"""
Diagnostic: test different CIQRANGE syntax variations for EV multiples.
"""
import time
import os
import pythoncom
from openpyxl import Workbook
from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

pythoncom.CoInitialize()

XLSX = os.path.join(r"C:\Claude", "ev_diag.xlsx")

# Test different formula syntaxes for IQ_TEV_EBITDA
FORMULAS = [
    # Row 2: Current single value (known working)
    ("CIQ single (current)", '=CIQ("DSGX","IQ_TEV_EBITDA")'),
    # Row 3: CIQRANGE multiples syntax: (Ticker, Metric, Period, D1, D2)
    ("CIQRANGE multi: FQ,D1,D2", '=CIQRANGE("DSGX","IQ_TEV_EBITDA",IQ_FQ,"02/26/2025","02/26/2026")'),
    # Row 4: CIQRANGE multiples syntax with FY
    ("CIQRANGE multi: FY,D1,D2", '=CIQRANGE("DSGX","IQ_TEV_EBITDA",IQ_FY,"02/26/2025","02/26/2026")'),
    # Row 5: CIQRANGE date-range only (market data style)
    ("CIQRANGE daterange D1,D2", '=CIQRANGE("DSGX","IQ_TEV_EBITDA","02/26/2025","02/26/2026")'),
    # Row 6: CIQRANGE financial style: IQ_FQ-4 (go back 4 quarters)
    ("CIQRANGE fin: IQ_FQ-4", '=CIQRANGE("DSGX","IQ_TEV_EBITDA",IQ_FQ-4)'),
    # Row 7: CIQRANGE financial style with label
    ("CIQRANGE fin: IQ_FQ-4 label", '=CIQRANGE("DSGX","IQ_TEV_EBITDA",IQ_FQ-4,,,,,,"TEV/EBITDA")'),
    # Row 8: TEV as market data (date range - should work)
    ("CIQRANGE TEV daterange", '=CIQRANGE("DSGX","IQ_TEV","02/26/2025","02/26/2026",,,,,,"TEV")'),
    # Row 9: EBITDA as financial item
    ("CIQRANGE EBITDA fin", '=CIQRANGE("DSGX","IQ_EBITDA",IQ_FQ-4,,,,,,"EBITDA")'),
    # Row 10: Revenue as financial item
    ("CIQRANGE Revenue fin", '=CIQRANGE("DSGX","IQ_TOTAL_REV",IQ_FQ-4,,,,,,"Revenue")'),
    # Row 11: Close price (known working market data)
    ("CIQRANGE Close price", '=CIQRANGE("DSGX","IQ_CLOSEPRICE","02/26/2025","02/26/2026",,,,,,"Price")'),
    # Row 12: Market cap as market data
    ("CIQRANGE MktCap daterange", '=CIQRANGE("DSGX","IQ_MARKETCAP","02/26/2025","02/26/2026",,,,,,"MktCap")'),
]

wb = Workbook()
ws = wb.active
ws.title = "Diag"
ws.cell(row=1, column=1, value="Test")
ws.cell(row=1, column=2, value="Formula")
ws.cell(row=1, column=3, value="Result (col B)")

for i, (desc, formula) in enumerate(FORMULAS):
    row = i + 2
    ws.cell(row=row, column=1, value=desc)
    ws.cell(row=row, column=2, value=formula)

wb.save(XLSX)
print(f"Created: {XLSX}")
print(f"Testing {len(FORMULAS)} formula variations\n")

# Open in Excel
session = launch_excel_isolated(XLSX)
excel = session.excel

try:
    excel.Run("SNLXLAddin.xla!RefreshSheet")
    print("Triggered refresh")
except Exception as e:
    print(f"Refresh: {e}")

# Wait for evaluation
print("Waiting 60s for evaluation...")
for tick in range(12):
    time.sleep(5)
    ws_com = session.workbook.Sheets(1)
    # Check a few cells
    val2 = ws_com.Cells(2, 2).Value
    val3 = ws_com.Cells(3, 2).Value
    elapsed = (tick + 1) * 5
    print(f"  [{elapsed}s] Row2={val2!r}, Row3={val3!r}")

    # If row 2 has a numeric value, data is flowing
    if isinstance(val2, (int, float)) and val2 > 0:
        time.sleep(15)  # let everything settle
        break

# Read all results
print("\n" + "=" * 80)
print("RESULTS:")
print("=" * 80)
ws_com = session.workbook.Sheets(1)

for i, (desc, formula) in enumerate(FORMULAS):
    row = i + 2
    val = ws_com.Cells(row, 2).Value

    # Also check if data expanded into more rows/columns
    extra_info = ""
    # Check column C and D for spill-over
    val_c = ws_com.Cells(row, 3).Value
    val_d = ws_com.Cells(row, 4).Value
    if val_c is not None:
        extra_info += f" | C={val_c!r}"
    if val_d is not None:
        extra_info += f" | D={val_d!r}"

    # Check if CIQRANGE expanded downward (check a few rows below)
    expanded_count = 0
    if "CIQRANGE" in desc.upper() or "RANGE" in desc.upper():
        for check_row in range(row + 1, row + 6):
            check_val = ws_com.Cells(check_row, 2).Value
            # Only count if it's a genuine expanded value (not another test formula)
            if check_val is not None and check_row <= len(FORMULAS) + 1:
                break  # Would overlap with next test row
            if check_val is not None:
                expanded_count += 1
    if expanded_count > 0:
        extra_info += f" | +{expanded_count} rows expanded"

    print(f"  [{row:2d}] {desc:<35} -> {val!r}{extra_info}")

close_session(session, save=False, delete_workbook=True)
print("\nDone.")
