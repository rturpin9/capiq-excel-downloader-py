"""
Diagnostic 3: Test the correct CIQRANGE layout (formula in row 1, date column).
Also test CIQ() with period parameter for pre-calculated multiples.
"""
import time
import os
import pythoncom
from openpyxl import Workbook
from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

pythoncom.CoInitialize()

XLSX = os.path.join(r"C:\Claude", "ev_diag3.xlsx")

wb = Workbook()
wb.remove(wb.active)

# Sheet 1: Correct layout — formulas in ROW 1, data expands downward
ws = wb.create_sheet(title="CorrectLayout")
# Column A: Date (financial period dates)
ws.cell(row=1, column=1, value='=CIQRANGE("DSGX","IQ_PERIODDATE_BS",IQ_FQ-12,,,,,,,"Date")')
# Column B: EBITDA
ws.cell(row=1, column=2, value='=CIQRANGE("DSGX","IQ_EBITDA",IQ_FQ-12,,,,,,,"EBITDA")')
# Column C: Revenue
ws.cell(row=1, column=3, value='=CIQRANGE("DSGX","IQ_TOTAL_REV",IQ_FQ-12,,,,,,,"Revenue")')
# Column D: TEV (try financial period syntax for market data)
ws.cell(row=1, column=4, value='=CIQRANGE("DSGX","IQ_TEV",IQ_FQ-12,,,,,,,"TEV")')
# Column E: Market Cap (financial period syntax)
ws.cell(row=1, column=5, value='=CIQRANGE("DSGX","IQ_MARKETCAP",IQ_FQ-12,,,,,,,"MktCap")')

# Sheet 2: CIQ() with period parameters for pre-calculated multiples
ws2 = wb.create_sheet(title="CIQ_Periods")
ws2.cell(row=1, column=1, value="Period")
ws2.cell(row=1, column=2, value="EV/EBITDA")
ws2.cell(row=1, column=3, value="EV/Revenue")
ws2.cell(row=1, column=4, value="Date")

for i in range(13):  # quarters 0 to 12
    row = i + 2
    period = f"IQ_FQ-{i}"
    ws2.cell(row=row, column=1, value=period)
    ws2.cell(row=row, column=2, value=f'=CIQ("DSGX","IQ_TEV_EBITDA","{period}")')
    ws2.cell(row=row, column=3, value=f'=CIQ("DSGX","IQ_TEV_TOTAL_REV","{period}")')
    ws2.cell(row=row, column=4, value=f'=CIQ("DSGX","IQ_PERIODDATE_BS","{period}")')

# Sheet 3: CIQ() with IQ_LTM, IQ_NTM etc for multiples
ws3 = wb.create_sheet(title="CIQ_LTM")
ws3.cell(row=1, column=1, value="Test")
ws3.cell(row=1, column=2, value="Value")
tests3 = [
    ('CIQ TEV_EBITDA noperiod', '=CIQ("DSGX","IQ_TEV_EBITDA")'),
    ('CIQ TEV_EBITDA IQ_LTM', '=CIQ("DSGX","IQ_TEV_EBITDA","IQ_LTM")'),
    ('CIQ TEV_EBITDA IQ_FQ-0', '=CIQ("DSGX","IQ_TEV_EBITDA","IQ_FQ-0")'),
    ('CIQ TEV_EBITDA IQ_FQ-1', '=CIQ("DSGX","IQ_TEV_EBITDA","IQ_FQ-1")'),
    ('CIQ TEV_EBITDA IQ_FQ-4', '=CIQ("DSGX","IQ_TEV_EBITDA","IQ_FQ-4")'),
    ('CIQ EBITDA IQ_FQ-0', '=CIQ("DSGX","IQ_EBITDA","IQ_FQ-0")'),
    ('CIQ EBITDA IQ_FQ-1', '=CIQ("DSGX","IQ_EBITDA","IQ_FQ-1")'),
]
for i, (desc, formula) in enumerate(tests3):
    row = i + 2
    ws3.cell(row=row, column=1, value=desc)
    ws3.cell(row=row, column=2, value=formula)

wb.save(XLSX)
print(f"Created: {XLSX}\n")

# Open in Excel
session = launch_excel_isolated(XLSX)
excel = session.excel

try:
    excel.Run("SNLXLAddin.xla!RefreshWorkbook")
    print("Triggered RefreshWorkbook")
except Exception as e:
    print(f"Refresh: {e}")

# Wait
print("Waiting for evaluation...")
for tick in range(24):
    time.sleep(5)
    elapsed = (tick + 1) * 5
    ws0 = session.workbook.Sheets(1)
    val_a2 = ws0.Cells(2, 1).Value
    val_b2 = ws0.Cells(2, 2).Value
    print(f"  [{elapsed}s] Sheet1: A2={val_a2!r} B2={val_b2!r}")
    if isinstance(val_b2, (int, float)) and val_b2 != 0:
        time.sleep(15)
        break

# Read results
print("\n" + "=" * 80)
print("Sheet 1: CorrectLayout (formulas in row 1)")
print("=" * 80)
ws = session.workbook.Sheets(1)
for row in range(1, 20):
    vals = []
    for col in range(1, 6):
        v = ws.Cells(row, col).Value
        if v is not None and hasattr(v, 'strftime'):
            v = v.strftime('%Y-%m-%d')
        vals.append(repr(v))
    print(f"  Row {row:2d}: {', '.join(vals)}")

print("\n" + "=" * 80)
print("Sheet 2: CIQ() with period parameters")
print("=" * 80)
ws2 = session.workbook.Sheets(2)
for row in range(1, 16):
    vals = []
    for col in range(1, 5):
        v = ws2.Cells(row, col).Value
        if v is not None and hasattr(v, 'strftime'):
            v = v.strftime('%Y-%m-%d')
        vals.append(repr(v))
    print(f"  Row {row:2d}: {', '.join(vals)}")

print("\n" + "=" * 80)
print("Sheet 3: CIQ() with different period specifiers")
print("=" * 80)
ws3 = session.workbook.Sheets(3)
for row in range(1, len(tests3) + 2):
    desc = ws3.Cells(row, 1).Value
    val = ws3.Cells(row, 2).Value
    print(f"  Row {row:2d}: {desc!r:40s} -> {val!r}")

close_session(session, save=False, delete_workbook=True)
print("\nDone.")
