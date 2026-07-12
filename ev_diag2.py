"""
Diagnostic 2: Test CIQRANGE expansion for EV multiples components.
Each formula gets its own sheet to avoid expansion collisions.
"""
import time
import os
import pythoncom
from openpyxl import Workbook
from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

pythoncom.CoInitialize()

XLSX = os.path.join(r"C:\Claude", "ev_diag2.xlsx")

# Each test gets its own sheet
TESTS = [
    # (sheet_name, formula, description)
    ("TEV_EBITDA_FQ12", '=CIQRANGE("DSGX","IQ_TEV_EBITDA",IQ_FQ-12,,,,,,,"EV_EBITDA")', "Pre-calc multiple, IQ_FQ-12"),
    ("TEV_EBITDA_FQ4", '=CIQRANGE("DSGX","IQ_TEV_EBITDA",IQ_FQ-4,,,,,,,"EV_EBITDA")', "Pre-calc multiple, IQ_FQ-4"),
    ("TEV_daterange", '=CIQRANGE("DSGX","IQ_TEV","02/26/2025","02/26/2026",,,,,"TEV")', "TEV as market data"),
    ("EBITDA_FQ12", '=CIQRANGE("DSGX","IQ_EBITDA",IQ_FQ-12,,,,,,,"EBITDA")', "EBITDA financial, IQ_FQ-12"),
    ("Revenue_FQ12", '=CIQRANGE("DSGX","IQ_TOTAL_REV",IQ_FQ-12,,,,,,,"Revenue")', "Revenue financial, IQ_FQ-12"),
    ("Close_daterange", '=CIQRANGE("DSGX","IQ_CLOSEPRICE","02/26/2025","02/26/2026",,,,,"Price")', "Close price (known working)"),
    ("MktCap_daterange", '=CIQRANGE("DSGX","IQ_MARKETCAP","02/26/2025","02/26/2026",,,,,"MktCap")', "Market cap date range"),
]

wb = Workbook()
wb.remove(wb.active)

for sheet_name, formula, desc in TESTS:
    ws = wb.create_sheet(title=sheet_name)
    ws.cell(row=1, column=1, value="Date")
    ws.cell(row=1, column=2, value=desc)
    ws.cell(row=2, column=2, value=formula)

wb.save(XLSX)
print(f"Created: {XLSX} with {len(TESTS)} sheets\n")

# Open in Excel
session = launch_excel_isolated(XLSX)
excel = session.excel

try:
    excel.Run("SNLXLAddin.xla!RefreshWorkbook")
    print("Triggered RefreshWorkbook")
except Exception as e:
    print(f"Refresh: {e}")

# Wait for data
print("Waiting for evaluation...")
for tick in range(24):
    time.sleep(5)
    elapsed = (tick + 1) * 5
    ws0 = session.workbook.Sheets(1)
    val = ws0.Cells(2, 2).Value
    print(f"  [{elapsed}s] Sheet1 B2 = {val!r}")
    if isinstance(val, (int, float)) and val > 0:
        time.sleep(20)  # let all sheets finish
        break

# Read results from all sheets
print("\n" + "=" * 80)
print("RESULTS:")
print("=" * 80)

for i, (sheet_name, formula, desc) in enumerate(TESTS):
    ws = session.workbook.Sheets(i + 1)
    print(f"\n--- Sheet: {sheet_name} ({desc}) ---")

    # Read up to 20 rows
    data_points = 0
    for row in range(2, 22):
        date_val = ws.Cells(row, 1).Value
        data_val = ws.Cells(row, 2).Value

        if date_val is None and data_val is None:
            if row > 3:
                break
            continue

        data_points += 1
        # Format date if COM datetime
        date_str = "None"
        if date_val is not None:
            if hasattr(date_val, 'strftime'):
                date_str = date_val.strftime('%Y-%m-%d')
            else:
                date_str = repr(date_val)

        print(f"  Row {row:2d}: date={date_str:>12}  val={data_val!r}")

    if data_points == 0:
        print(f"  (no data in rows 2-21)")
    else:
        print(f"  Total: {data_points} data points")

close_session(session, save=False, delete_workbook=True)
print("\nDone.")
