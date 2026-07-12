"""
Diagnostic 4: Test correct multiples syntax with IQ_LTM + date range.
Both CIQRANGE and CIQRANGEV variants.
"""
import time
import os
import pythoncom
from openpyxl import Workbook
from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session

pythoncom.CoInitialize()

XLSX = os.path.join(r"C:\Claude", "ev_diag4.xlsx")

wb = Workbook()
wb.remove(wb.active)

TESTS = [
    # (sheet_name, formula, description)
    # CIQRANGEV with IQ_LTM + absolute dates
    ("V_LTM_abs", '=CIQRANGEV("DSGX","IQ_TEV_EBITDA",IQ_LTM,"2/26/2025","2/26/2026",,,,"TEV/EBITDA")', "CIQRANGEV IQ_LTM abs dates"),
    # CIQRANGEV with IQ_LTM + relative date
    ("V_LTM_rel", '=CIQRANGEV("DSGX","IQ_TEV_EBITDA",IQ_LTM,"-1Y","2/26/2026",,,,"TEV/EBITDA")', "CIQRANGEV IQ_LTM rel date"),
    # CIQRANGE (formulas) with IQ_LTM + absolute dates
    ("R_LTM_abs", '=CIQRANGE("DSGX","IQ_TEV_EBITDA",IQ_LTM,"2/26/2025","2/26/2026",,,,"TEV/EBITDA")', "CIQRANGE IQ_LTM abs dates"),
    # CIQRANGEV for EV/Revenue too
    ("V_Rev_abs", '=CIQRANGEV("DSGX","IQ_TEV_TOTAL_REV",IQ_LTM,"2/26/2025","2/26/2026",,,,"TEV/Revenue")', "CIQRANGEV EV/Rev abs dates"),
    # CIQRANGEV with 3-year lookback
    ("V_LTM_3Y", '=CIQRANGEV("DSGX","IQ_TEV_EBITDA",IQ_LTM,"-3Y","2/26/2026",,,,"TEV/EBITDA")', "CIQRANGEV IQ_LTM 3yr"),
]

for sheet_name, formula, desc in TESTS:
    ws = wb.create_sheet(title=sheet_name)
    ws.cell(row=1, column=1, value=formula)

wb.save(XLSX)
print(f"Created: {XLSX} with {len(TESTS)} sheets\n")

session = launch_excel_isolated(XLSX)
excel = session.excel

try:
    excel.Run("SNLXLAddin.xla!RefreshWorkbook")
    print("Triggered RefreshWorkbook")
except Exception as e:
    print(f"Refresh: {e}")

print("Waiting for evaluation...")
for tick in range(24):
    time.sleep(5)
    elapsed = (tick + 1) * 5
    ws0 = session.workbook.Sheets(1)
    val = ws0.Cells(1, 1).Value
    val2 = ws0.Cells(2, 1).Value
    print(f"  [{elapsed}s] A1={val!r}  A2={val2!r}")
    if isinstance(val2, (int, float)) and val2 > 0:
        time.sleep(15)
        break
    if isinstance(val, (int, float)) and val > 0 and tick >= 3:
        time.sleep(15)
        break

print("\n" + "=" * 80)
print("RESULTS:")
print("=" * 80)

for i, (sheet_name, formula, desc) in enumerate(TESTS):
    ws = session.workbook.Sheets(i + 1)
    print(f"\n--- {sheet_name}: {desc} ---")
    for row in range(1, 25):
        val = ws.Cells(row, 1).Value
        if val is None and row > 3:
            break
        if val is not None:
            if hasattr(val, 'strftime'):
                val = val.strftime('%Y-%m-%d')
            print(f"  Row {row:2d}: {val!r}")
    print()

close_session(session, save=False, delete_workbook=True)
print("Done.")
