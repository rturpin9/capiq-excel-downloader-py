"""Regression tests for workbook batching and dialect-specific lookup layout."""
import math

import pandas as pd
from openpyxl import Workbook, load_workbook

from capiq_excel.formulas.ciq_builder import CiqBuilder
from capiq_excel.formulas.spg_builder import SpgBuilder
from capiq_excel.config import FormulaOptions
from capiq_excel.workbook import create


def _workbook_and_sheet():
    workbook = Workbook()
    return workbook, workbook.active


def test_id_workbook_saves_final_partial_chunk(tmp_path, monkeypatch):
    monkeypatch.setattr(create, "get_workbook_and_worksheet", _workbook_and_sheet)
    identifiers = [f"T{i}" for i in range(101)]
    create.create_all_xlsx_with_id_commands(
        identifiers,
        str(tmp_path),
        num_files=math.ceil(len(identifiers) / 100),
        builder=CiqBuilder(),
    )

    files = sorted(tmp_path.glob("ids *.xlsx"))
    assert len(files) == 2
    row_counts = [load_workbook(path).active.max_row - 1 for path in files]
    assert row_counts == [51, 50]
    assert sum(row_counts) == len(identifiers)


def test_non_spill_lookup_values_use_retained_columns():
    frame = pd.DataFrame({"ID": ["MSFT"]})
    builder = SpgBuilder()
    create._fill_capiq_id_column(frame, builder)
    create._fill_capiq_name_column(frame, builder)

    assert list(frame.columns) == ["ID", "IQID", "IQ Name"]
    assert frame.loc[0, "IQID"].startswith("=SPG(")
    assert frame.loc[0, "IQ Name"].startswith("=SPG(")


def test_ciq_lookup_keeps_horizontal_spill_columns():
    frame = pd.DataFrame({"ID": ["MSFT"]})
    builder = CiqBuilder()
    create._fill_capiq_id_column(frame, builder)
    create._fill_capiq_name_column(frame, builder)

    assert list(frame.columns) == ["ID", "Blank 1", "IQID", "Blank 2", "IQ Name"]


def test_formula_options_reach_builder_generated_workbook():
    workbook = Workbook()
    sheet = workbook.active
    create._fill_with_commands(
        sheet,
        "MSFT",
        {"IQ_TOTAL_REV": "Revenue"},
        {},
        builder=SpgBuilder(),
        formula_options=FormulaOptions(currency="USD", magnitude="Millions"),
        freq="Q",
        num_periods=4,
    )
    assert "Curr=USD,Mag=Millions" in sheet["A1"].value
    assert "Curr=USD,Mag=Millions" in sheet["C1"].value
