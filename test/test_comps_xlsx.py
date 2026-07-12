"""Unit tests for the formatted comp-table XLSX writer.

Loads the produced workbook with openpyxl (no Excel/COM) and asserts on
structure, headers, number formats, derived formulas, and aggregation ranges.
"""
import pytest
from openpyxl import load_workbook

from capiq_excel.engines.comps import resolve_extras
from capiq_excel.engines.comps_xlsx import (
    build_columns, list_columns, write_comps_xlsx,
)


def _company(name, ticker, **kw):
    base = {
        "Company": name, "Ticker": ticker, "GAAP": "US GAAP",
        "Mkt Cap": 1000, "Total Debt": 200, "Leases": 50, "Cash": 100,
        "Net Debt": 100, "TEV": 1100, "LTM Rev": 800, "NTM Rev": 850,
        "LTM EBITDA": 200, "Lease Adj": 10, "NTM EBITDA": 220,
        "Lease Adj Applied": True, "EBITDA Margin %": 25.0, "Gross Margin %": 40.0,
        "CapEx": 30, "EV/Rev": 1.38, "EV/LTM EBITDA": 5.5, "EV/NTM EBITDA": 5.0,
        "P/BV": 2.0, "ND/EBITDA": 0.5,
    }
    base.update(kw)
    return base


def _data(mode="lease-adjusted", companies=None):
    return {
        "mode": mode, "currency": "CAD", "date": "2/26/2026",
        "lease_adjust_ntm": True,
        "companies": companies if companies is not None else [
            _company("Alpha Inc", "AAA"), _company("Beta Corp", "BBB"),
        ],
    }


def _label_rows(ws):
    """Map the first-table-column (C) label text -> row number."""
    out = {}
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row, 3).value
        if isinstance(v, str) and v:
            out[v] = row
    return out


# ── build_columns ──────────────────────────────────────────────────────────

class TestBuildColumns:
    def test_default_lease_adjusted_order(self):
        keys = [c.key for c in build_columns("lease-adjusted", None, None)]
        assert keys[0] == "company"
        assert "lease_adj" in keys
        assert keys[-1] == "nd_ebitda"

    def test_excluding_leases_drops_lease_adj(self):
        keys = [c.key for c in build_columns("excluding-leases", None, None)]
        assert "lease_adj" not in keys
        assert "tev" in keys

    def test_custom_selection_prepends_company(self):
        cols = build_columns("lease-adjusted", ["mkt_cap", "tev", "ev_rev"], None)
        keys = [c.key for c in cols]
        assert keys[0] == "company"          # auto-prepended
        assert keys[1:] == ["mkt_cap", "tev", "ev_rev"]

    def test_unknown_key_raises(self):
        with pytest.raises(ValueError, match="Unknown column"):
            build_columns("lease-adjusted", ["company", "bogus"], None)

    def test_mode_inapplicable_key_skipped_not_raised(self):
        # lease_adj is valid in lease-adjusted mode but absent in excluding-leases;
        # selecting it for the excl tab should skip it, not error.
        keys = [c.key for c in
                build_columns("excluding-leases", ["company", "lease_adj", "tev"], None)]
        assert keys == ["company", "tev"]

    def test_extras_appended_and_selectable(self):
        specs, _ = resolve_extras(["roe"])
        keys = [c.key for c in build_columns("lease-adjusted", None, specs)]
        assert keys[-1] == "roe"
        # extra key is also selectable explicitly
        keys2 = [c.key for c in build_columns("lease-adjusted", ["company", "roe"], specs)]
        assert keys2 == ["company", "roe"]


class TestListColumns:
    def test_defaults_marked(self):
        rows = list_columns("lease-adjusted")
        defaults = {k for k, _, is_def in rows if is_def}
        assert {"company", "mkt_cap", "tev", "ev_ltm_ebitda"} <= defaults
        assert "capex" not in defaults     # available but not default


# ── write_comps_xlsx ─────────────────────────────────────────────────────────

class TestWriteXlsx:
    def test_titles_and_headers(self, tmp_path):
        p = tmp_path / "c.xlsx"
        write_comps_xlsx(_data(), str(p))
        ws = load_workbook(p).active
        assert ws.title == "Lease Adjusted"
        assert ws.cell(2, 3).value == "Comparable Companies Analysis"
        assert "As at February 26, 2026" in ws.cell(3, 3).value
        assert "CAD$ millions" in ws.cell(3, 3).value
        headers = [ws.cell(4, ci).value for ci in range(3, 3 + 18)]
        assert headers[0] == "Company"
        assert "Enterprise\nValue" in headers

    def test_sheet_name_excluding(self, tmp_path):
        p = tmp_path / "c.xlsx"
        d = _data(mode="excluding-leases")
        d.pop("lease_adjust_ntm", None)
        write_comps_xlsx(d, str(p))
        assert load_workbook(p).active.title == "Excluding Leases"

    def test_derived_formulas_and_value_columns(self, tmp_path):
        p = tmp_path / "c.xlsx"
        write_comps_xlsx(_data(), str(p))
        ws = load_workbook(p).active
        rows = _label_rows(ws)
        row = rows["Alpha Inc *"]   # name carries lease-adj asterisk
        # EBITDA margin (col O) and EV multiples (P/Q/R) and ND/EBITDA (T) = formulas
        assert ws.cell(row, 15).value.startswith("=IFERROR(")     # O ebitda_margin
        assert "/J" in ws.cell(row, 15).value
        assert ws.cell(row, 16).value.startswith("=IFERROR(IF(")  # P ev_rev
        # P/BV (col S) is a static value, not a formula
        assert ws.cell(row, 19).value == 2.0
        # money value (Market Cap col D)
        assert ws.cell(row, 4).value == 1000

    def test_ticker_and_name_fallback(self, tmp_path):
        p = tmp_path / "c.xlsx"
        d = _data(companies=[
            _company("Good Co", "GOOD"),
            _company("#INVALID COMPANY ID", "BADTKR"),  # unresolved name
        ])
        write_comps_xlsx(d, str(p))
        ws = load_workbook(p).active
        labels = _label_rows(ws)
        # bad name falls back to the ticker
        assert "BADTKR *" in labels
        assert "#INVALID COMPANY ID" not in labels
        # ticker floats in column B (green), for the good company
        good_row = labels["Good Co *"]
        assert ws.cell(good_row, 2).value == "GOOD"

    def test_missing_value_is_blank(self, tmp_path):
        p = tmp_path / "c.xlsx"
        c = _company("Sparse Co", "SPRS")
        del c["P/BV"]          # NaN-stripped key absent
        write_comps_xlsx(_data(companies=[c]), str(p))
        ws = load_workbook(p).active
        row = _label_rows(ws)["Sparse Co *"]
        assert ws.cell(row, 19).value is None   # S = P/BV blank

    def test_grouped_overall_excludes_group_stat_rows(self, tmp_path):
        """Overall Average must aggregate company rows only (regression)."""
        p = tmp_path / "c.xlsx"
        write_comps_xlsx(
            _data(companies=[_company("Alpha", "AAA"), _company("Beta", "BBB")]),
            str(p),
            groups={"G1": ["AAA"], "G2": ["BBB"]},
        )
        ws = load_workbook(p).active
        rows = _label_rows(ws)
        overall = rows["Overall Average"]
        f = ws.cell(overall, 15).value   # O column = ebitda_margin
        assert f.startswith("=IFERROR(AVERAGE(")
        # references both single-company rows, comma-joined…
        assert "," in f
        # …and must NOT reach into a per-group Average/Median row
        avg_g1 = rows["Average G1"]
        assert f"O{avg_g1}" not in f

    def test_number_formats(self, tmp_path):
        p = tmp_path / "c.xlsx"
        write_comps_xlsx(_data(), str(p))
        ws = load_workbook(p).active
        row = _label_rows(ws)["Alpha Inc *"]
        assert ws.cell(row, 4).number_format.startswith("_(*")   # money accounting
        assert ws.cell(row, 15).number_format == r'0.0%;(0.0%)'  # pct
        assert ws.cell(row, 16).number_format == r'#,##0.0"x";(#,##0.0"x")'  # mult

    def test_negative_net_debt_multiple_is_preserved(self, tmp_path):
        p = tmp_path / "negative-nd.xlsx"
        company = _company("Net Cash Co", "CASH", **{
            "Net Debt": -100,
            "ND/EBITDA": -0.5,
        })
        write_comps_xlsx(_data(companies=[company]), str(p))
        ws = load_workbook(p).active
        row = _label_rows(ws)["Net Cash Co *"]
        formula = ws.cell(row, 20).value
        assert formula.startswith("=IFERROR(")
        assert ">0" not in formula

        p2 = tmp_path / "negative-nd-static.xlsx"
        write_comps_xlsx(
            _data(companies=[company]),
            str(p2),
            columns=["company", "nd_ebitda"],
        )
        ws2 = load_workbook(p2).active
        row2 = _label_rows(ws2)["Net Cash Co *"]
        assert ws2.cell(row2, 4).value == -0.5


class TestMultiTab:
    def test_two_tabs_named_and_ordered(self, tmp_path):
        p = tmp_path / "c.xlsx"
        lease = _data(mode="lease-adjusted")
        excl = _data(mode="excluding-leases")
        excl.pop("lease_adjust_ntm", None)
        write_comps_xlsx([lease, excl], str(p), active_mode="excluding-leases")
        wb = load_workbook(p)
        assert wb.sheetnames == ["Lease Adjusted", "Excluding Leases"]
        # excluding-leases tab must not carry the lease_adj column header
        excl_ws = wb["Excluding Leases"]
        headers = [excl_ws.cell(4, ci).value for ci in range(3, 25)]
        assert "Lease\nAdj." not in headers
        # active tab honored
        assert wb.active.title == "Excluding Leases"

    def test_single_dict_still_one_tab(self, tmp_path):
        p = tmp_path / "c.xlsx"
        write_comps_xlsx(_data(), str(p))
        assert load_workbook(p).sheetnames == ["Lease Adjusted"]
