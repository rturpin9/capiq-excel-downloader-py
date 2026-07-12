"""
Test identifier lookup via CIQ() formulas.

Resolves a mix of identifier types to Capital IQ IDs and company names.
"""
import os
import time
import openpyxl
import pythoncom

from capiq_excel.excel_lifecycle import launch_excel_isolated, close_session
from capiq_excel.formulas.ciq_builder import CiqBuilder

pythoncom.CoInitialize()

# Mix of identifier types to test
TEST_IDS = [
    "NVDA",                # US ticker
    "CRM",                 # US ticker
    "TSX:CSU",             # Canadian ticker with exchange prefix
    "Microsoft",           # Company name
    "03783310",            # CUSIP (Apple)
    "DSGX",               # US ticker
    "CMDXF",              # US OTC ticker
]

BASE_DIR = r"C:\Claude"
XLSX_PATH = os.path.join(BASE_DIR, "test_id_lookup.xlsx")

builder = CiqBuilder()


def create_workbook():
    """Create XLSX with CIQ ID and name lookup formulas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ID Lookup"

    # Headers
    ws.cell(row=1, column=1, value="Input")
    ws.cell(row=1, column=2, value="CIQ ID")
    ws.cell(row=1, column=3, value="Company Name")

    for i, identifier in enumerate(TEST_IDS):
        row = i + 2
        ws.cell(row=row, column=1, value=identifier)
        ws.cell(row=row, column=2,
                value=builder.build_identifier_lookup(identifier, "IQ_COMPANY_ID_QUICK_MATCH"))
        ws.cell(row=row, column=3,
                value=builder.build_identifier_lookup(identifier, "IQ_COMPANY_NAME_QUICK_MATCH"))

    wb.save(XLSX_PATH)
    print(f"Created workbook: {XLSX_PATH}")

    # Show the formulas
    print("\nFormulas written:")
    for i, identifier in enumerate(TEST_IDS):
        id_formula = builder.build_identifier_lookup(identifier, "IQ_COMPANY_ID_QUICK_MATCH")
        name_formula = builder.build_identifier_lookup(identifier, "IQ_COMPANY_NAME_QUICK_MATCH")
        print(f"  {identifier:>15}  ->  {id_formula}  |  {name_formula}")


def open_refresh_and_extract():
    """Launch Excel, refresh, extract resolved IDs."""
    print("\nLaunching Excel...")
    session = launch_excel_isolated(XLSX_PATH)
    excel = session.excel
    print(f"Connected to Excel {excel.Version}")

    try:
        excel.Run("SNLXLAddin.xla!RefreshWorkbook")
        print("Triggered RefreshWorkbook")
    except Exception as e:
        print(f"Refresh warning: {e}")

    # Poll for completion
    print("Waiting for CIQ to evaluate...")
    for tick in range(24):
        time.sleep(5)
        elapsed = (tick + 1) * 5
        ws = session.workbook.Sheets(1)
        val = ws.Cells(2, 2).Value  # First CIQ ID result
        print(f"  [{elapsed}s] B2={val!r}")
        if val is not None and isinstance(val, (int, float, str)) and str(val).startswith("IQ"):
            time.sleep(5)  # let remaining cells settle
            break
        if isinstance(val, (int, float)) and val > 0:
            time.sleep(5)
            break
    else:
        print("  Timeout — extracting whatever is available")

    # Extract results
    ws = session.workbook.Sheets(1)
    print("\n" + "=" * 80)
    print(f"{'Input':<20} {'CIQ ID':<20} {'Company Name':<40}")
    print("-" * 80)

    results = []
    for i in range(len(TEST_IDS)):
        row = i + 2
        input_id = ws.Cells(row, 1).Value
        ciq_id = ws.Cells(row, 2).Value
        name = ws.Cells(row, 3).Value

        # Convert COM values to strings
        ciq_id_str = str(ciq_id) if ciq_id is not None else "(empty)"
        name_str = str(name) if name is not None else "(empty)"

        print(f"{str(input_id):<20} {ciq_id_str:<20} {name_str:<40}")
        results.append((input_id, ciq_id, name))

    print("=" * 80)

    # Summary
    resolved = sum(1 for _, cid, _ in results if cid is not None and str(cid).startswith("IQ"))
    failed = sum(1 for _, cid, _ in results
                 if cid is not None and isinstance(cid, str)
                 and any(tok in cid.upper() for tok in ("#", "ERROR", "INVALID", "NA")))
    pending = sum(1 for _, cid, _ in results if cid is None or cid == "")

    print(f"\nResolved: {resolved}/{len(TEST_IDS)}")
    if failed:
        print(f"Failed:   {failed}/{len(TEST_IDS)}")
    if pending:
        print(f"Pending:  {pending}/{len(TEST_IDS)}")

    close_session(session, save=False, delete_workbook=False)
    return results


if __name__ == "__main__":
    print("=" * 50)
    print("  CIQ Identifier Lookup Test")
    print(f"  Testing {len(TEST_IDS)} identifiers")
    print("=" * 50)

    create_workbook()
    results = open_refresh_and_extract()
